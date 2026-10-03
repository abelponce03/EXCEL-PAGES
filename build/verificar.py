"""Verifica el libro recalculado contra un cálculo independiente en Python."""
import math, sys
from openpyxl import load_workbook
f = sys.argv[1]
wb = load_workbook(f, data_only=True)
ca = wb["CALCULO"]
hdr = {ca.cell(5, c).value: c for c in range(1, ca.max_column + 1)}
def v(r, h): return ca.cell(r, hdr[h]).value

# ---- cálculo independiente (parámetros de ejemplo)
P = dict(t=0.80, diesel=2.0, cita=10, dieta=10, jor=15, jorct=15, com=0.03, desg=0.12)
rev = 0.005 + 0.01 + 0.10 + 0.01
conts = [  # kgO,kgC,kgR, litP, jorT, estad, litDO, litTC, litDC, jorC, litTR, litDR, jorR, otrG, otrI, mes
 (5200,2600,2300, 62, 11, 0, None,None,None,None,None,None,None, 0, 0, 9),
 (4600,2100,2700, None,None,120, None,None,None,None,480,None,None, 0, 0, 9),
 (6100,2900,2400, None,None,0, 150,None,None,None,None,260,None, 35, 150, 10),
 (3800,1500,3900, None,None,0, None,None,None,None,None,None,None, 0, 0, 10)]
fixG = 60+250+21*50+200+40+30+15+50+10 + (45000-5000)/8/12 + 1500/5/12 + 150 + 1020/12 + 0
fixG += (60+250+21*50)*0.19
fixR = 200
res = []
for c in conts:
    kO,kC,kR,litP,jorT,est,lDO,lTC,lDC,jC,lTR,lDR,jR,oG,oI,m = c
    kT = kO+kC+kR
    ing = kT*P['t'] + oI
    litPs = 60; litP = litP if litP is not None else litPs
    puerto = P['cita'] + litP*2 + P['dieta'] + 180*P['desg']
    jT = jorT if jorT is not None else math.ceil(kT/1000)
    trans = jT*P['jor'] + est
    vO = math.ceil(kO/5000); vC = math.ceil(kC/5000); vR = math.ceil(kR/5000)
    dO = (lDO if lDO is not None else vO*120)*2
    dC = ((lTC or vC*250) + (lDC or vC*150))*2 + (jC or math.ceil(kC/800))*P['jorct']
    dR = ((lTR or vR*450) + (lDR or vR*180))*2 + (jR or math.ceil(kR/800))*P['jorct']
    com = kT*P['t']*P['com']
    var = puerto+trans+dO+dC+dR+com+ing*(0.005+0.01)+oG
    res.append(dict(kT=kT,kO=kO,kC=kC,kR=kR,ing=ing,var=var,m=m,puerto=puerto,trans=trans,dist=dO+dC+dR))
for m in (9,10):
    ms = [x for x in res if x['m']==m]
    kgm = sum(x['kT'] for x in ms)
    for x in ms:
        fx = fixG*x['kT']/kgm + sum(fixR*x[k]/sum(y[k] for y in ms) for k in ('kO','kC','kR'))
        x['fix'] = fx
        x['uai'] = x['ing'] - x['var'] - fx - x['ing']*0.11
ok = True
for i, x in enumerate(res):
    r = 6+i
    pairs = [("INGRESO TOTAL", x['ing']), ("TOTAL PUERTO", x['puerto']), ("TOTAL TRANSITARIA", x['trans']),
             ("TOTAL DISTRIBUCIÓN", x['dist']), ("TOTAL COSTOS VARIABLES", x['var']),
             ("TOTAL FIJOS ASIGNADOS", x['fix']), ("UTILIDAD ANTES DE IMPUESTO SOBRE UTILIDADES", x['uai'])]
    for h, e in pairs:
        got = v(r, h)
        flag = abs(got-e) < 0.01
        ok &= flag
        print(f"cont{i+1} {h[:28]:28s} excel={got:12.2f} python={e:12.2f} {'OK' if flag else 'DIFERENCIA'}")
    print(f"   regiones: OCC={v(r,'Utilidad Occidente'):.2f} CEN={v(r,'Utilidad Centro'):.2f} ORI={v(r,'Utilidad Oriente'):.2f} "
          f"cuadre={v(r,'Control: diferencia de cuadre (debe ser 0)')} alerta='{v(r,'ALERTAS')}' desvío={v(r,'Desvío del combustible frente a la norma'):.3f}")
rm = wb["RESUMEN_MES"]
print("\nRESUMEN_MES")
for r in range(6, 9):
    print([rm.cell(r, c).value for c in (1,2,6,7,8,9,11,12,13,14,15,20,21,22,27)])
print("total", [rm.cell(30, c).value for c in (2,6,7,15,27)])
for r in range(33, 37):
    print([rm.cell(r, c).value for c in range(1, 6)])
fixAll = fixG + 3*fixR
print("fijos mes python:", round(fixAll,2), " uai sep python:", round(sum(x['ing']-x['var']-x['ing']*0.11 for x in res if x['m']==9)-fixAll,2))
db = wb["DASHBOARD"]
print("\nDASHBOARD")
for r in range(8, 19):
    print(f"  {db.cell(r,2).value:40s} {db.cell(r,3).value!s:>14}   {db.cell(r,7).value!s:42s} {db.cell(r,8).value}")
for r in range(21, 33):
    print("  ", [db.cell(r, c).value for c in range(2, 6)], [db.cell(r, c).value for c in range(7, 12)])
sm = wb["SIMULADOR"]
print("\nSIMULADOR")
for r in range(5, sm.max_row+1):
    vals = [sm.cell(r, c).value for c in range(2, 10)]
    if any(x is not None for x in vals): print(r, vals)
co = wb["COMERCIAL"]
print("\nCOMERCIAL", [co.cell(r, 7).value for r in range(35, 41)], co["C7"].value, co["I33"].value)
ini = wb["INICIO"]
for r in range(29, 37): print("INICIO", ini.cell(r,2).value, ini.cell(r,3).value)
cf = wb["COSTOS_FIJOS"]
for r in range(6, cf.max_row+1):
    if cf.cell(r,2).value: print("CF", r, cf.cell(r,2).value[:45], cf.cell(r,6).value, cf.cell(r,7).value, cf.cell(r,8).value)
print("\nTODO OK" if ok else "\nHAY DIFERENCIAS")
