# -*- coding: utf-8 -*-
"""Genera las capturas de pantalla y los valores esperados de la guía en PDF.

Para cada ejercicio de la guía se crea una copia real del libro con los datos del
ejercicio, se recalcula con LibreOffice y se fotografían los rangos indicados
(con letras de columna y números de fila, como en Excel). Los recuadros rojos
numerados marcan dónde debe mirar o escribir el usuario.

Salidas:
  guia/img/*.png    capturas
  guia/valores.tex  macros LaTeX con los resultados que el usuario debe ver

Uso: RECALC=/ruta/recalc.py python build/capturas.py [libro.xlsx]
Requiere LibreOffice Calc (soffice), pdftoppm (poppler) y Pillow.
"""
import concurrent.futures as cf
import datetime as dt
import json
import os
import shutil
import subprocess
import sys
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
LIBRO = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Analisis_Rentabilidad_Paqueteria.xlsx"
IMG = ROOT / "guia" / "img"
VALORES = ROOT / "guia" / "valores.tex"
RECALC = os.environ.get("RECALC")
TMP = Path(tempfile.mkdtemp(prefix="capturas_"))
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
D = dt.datetime


# ------------------------------------------------------------------ utilidades
def find_row(ws, col, text, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws[f"{col}{r}"].value
        if isinstance(v, str) and v.strip().startswith(text):
            return r
    raise KeyError(f"{ws.title}: no encuentro «{text}» en la columna {col}")


def col_of(ws, header, row=5):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row, c).value
        if isinstance(v, str) and v.strip() == header:
            return get_column_letter(c)
    raise KeyError(f"{ws.title}: no encuentro el encabezado «{header}»")


# ------------------------------------------------------------------ estados (ejercicios)
def s_nuevo(wb):
    """Ejercicio 2: registrar un contenedor nuevo al llegar."""
    e = wb["ENTRADA"]
    for c, v in zip("BCDEFGHI", ["CONT-PRUEBA-05", D(2026, 10, 27), "ENCI", "40' HC", 1000, 5000, 2500, 2500]):
        e[f"{c}10"] = v
    e["Z10"] = "Abierto"


def s_reales(wb):
    """Ejercicio 3: completar los datos reales y cerrar."""
    e = wb["ENTRADA"]
    e["K10"] = 64
    e["M10"] = 11
    e["S10"] = 520
    e["Y10"] = D(2026, 11, 6)
    e["Z10"] = "Cerrado"
    e["AA10"] = "Vales 1520-1528; nómina de desagrupe 45"


def s_precio(wb):
    """Ejercicio 4: sube el diésel desde el 1 de noviembre y llega un contenedor en noviembre."""
    p = wb["PRECIOS"]
    for c in "BCDEFGHIJKLMN":
        p[f"{c}7"] = p[f"{c}6"].value
    p["A7"] = D(2026, 11, 1)
    p["E7"] = 2.30
    p["O7"] = "EJERCICIO: aumento del diésel (factura CUPET 0457)"
    e = wb["ENTRADA"]
    for c, v in zip("BCDEFGHI", ["CONT-PRUEBA-06", D(2026, 11, 10), "A.V", "40'", 820, 4000, 2000, 2000]):
        e[f"{c}11"] = v
    e["Z11"] = "Abierto"


def s_cierre(wb):
    """Ejercicio 5: cierre de octubre (gasto real de electricidad y liquidación comercial)."""
    cfs = wb["COSTOS_FIJOS"]
    r = find_row(cfs, "B", "Electricidad")
    cfs[f"H{r}"] = 55
    co = wb["COMERCIAL"]
    extra = {"La Habana": 4000, "Villa Clara": 2000, "Santiago de Cuba": 2000}
    for rr in range(12, 33):
        t = co[f"D{rr}"].value
        if t in extra:
            co[f"G{rr}"] = round((co[f"G{rr}"].value or 0) + extra[t], 2)


def s_dashboard(wb):
    d = wb["DASHBOARD"]
    d["C4"] = D(2026, 10, 1)
    d["C5"] = D(2026, 11, 1)


def s_simula(wb):
    wb["SIMULADOR"]["D11"] = 0.95


def s_vacio(wb):
    """Ejercicio 8: borrar los contenedores de ejemplo y la base variable de COMERCIAL."""
    e = wb["ENTRADA"]
    for r in range(6, 10):
        for c in range(2, 28):
            e.cell(r, c).value = None
    co = wb["COMERCIAL"]
    for r in range(12, 33):
        co[f"G{r}"] = None


ESTADOS = {
    "s0": [],
    "s1": [s_nuevo],
    "s2": [s_nuevo, s_reales],
    "s3": [s_nuevo, s_reales, s_precio],
    "s4": [s_nuevo, s_reales, s_precio, s_cierre],
    "s5": [s_nuevo, s_reales, s_precio, s_cierre, s_dashboard],
    "s6": [s_simula],
    "s7": [s_vacio],
}


# ------------------------------------------------------------------ capturas
def calc_cols(*headers):
    """Rango de CALCULO que muestra solo las columnas indicadas (las demás se ocultan)."""
    return ("CALCULO", headers)


CAPTURAS = [
    # nombre, estado, hoja, rango, [recuadros numerados]
    ("inicio_indice", "s0", "INICIO", "B2:D23", ["B10:B22", "D10:D22"]),
    ("inicio_rutina", "s0", "INICIO", "B24:D34", ["B25:D28"]),
    ("inicio_controles", "s0", "INICIO", "B35:C45", ["C36:C42"]),
    ("entrada_izq", "s0", "ENTRADA", "A3:L11", ["B6:B9", "C6:E9", "G6:I9", "J6:L9"]),
    ("entrada_der", "s0", "ENTRADA", "V3:AC11", ["Y6:Z9", "AB6:AC9"]),
    ("parametros", "s0", "PARAMETROS", ("A4:E34", "B,C,D"), ["C6:C11", "C13:C25"]),
    ("precios", "s0", "PRECIOS", "A5:I8", ["A6:A6", "B6:I6"]),
    ("precios_b", "s0", "PRECIOS", "J5:O8", ["J6:N6", "O6:O6"]),
    ("activos", "s0", "ACTIVOS", "A5:J9", ["B6:G7", "H6:J7"]),
    ("activos_b", "s0", "ACTIVOS", "K5:O9", ["K6:M7"]),
    ("costos_fijos", "s0", "COSTOS_FIJOS", ("A5:I28", "A,B,C,E,F,G,H,I"), ["E6:E27", "F6:F27", "G6:I27"]),
    ("costos_fijos_tot", "s0", "COSTOS_FIJOS", "A28:I42", ["B29:B36", "G38:I42"]),
    ("comercial", "s0", "COMERCIAL", "A3:I20", ["C4:C4", "F12:G20"]),
    ("comercial_conc", "s0", "COMERCIAL", "A31:I41", ["G36:G39"]),
    ("calculo", "s0", "CALCULO", ("ID contenedor", "Fecha de arribo", "Mes", "kg total", "INGRESO TOTAL",
                                  "TOTAL PUERTO", "TOTAL TRANSITARIA", "TOTAL DISTRIBUCIÓN", "TOTAL COSTOS VARIABLES",
                                  "TOTAL FIJOS ASIGNADOS", "UTILIDAD ANTES DE IMPUESTO SOBRE UTILIDADES"), ["8:8"]),
    ("calculo_region", "s0", "CALCULO", ("ID contenedor", "Utilidad Occidente", "Utilidad Centro", "Utilidad Oriente",
                                         "Control: diferencia de cuadre (debe ser 0)", "ALERTAS"), ["8:8"]),
    ("resumen", "s0", "RESUMEN_MES", ("A3:P9", "A,B,C,F,G,H,I,K,M,N,O"), ["A6:A7", "O6:O7"]),
    ("resumen_pe", "s0", "RESUMEN_MES", ("A5:V9", "A,Q,R,S,T,U,V"), ["T6:V7"]),
    ("dashboard_kpi", "s0", "DASHBOARD", ("A1:H19", "B,C,G,H"), ["C4:C5", "C16:C18"]),
    ("dashboard_tablas", "s0", "DASHBOARD", "A20:L34", ["G22:K24"]),
    ("dashboard_graficos", "s0", "DASHBOARD", "A34:L72", []),
    ("simulador_sup", "s0", "SIMULADOR", "B3:F27", ["C5:C26", "D5:D26", "E5:E26"]),
    ("simulador_res", "s0", "SIMULADOR", "B29:F48", ["C38:F38", "C41:F41", "F48:F48"]),
    ("simulador_sens", "s0", "SIMULADOR", "B49:I60", ["E56:E56"]),
    ("simulador_vol", "s0", "SIMULADOR", "B67:D79", []),
    # ejercicios
    ("ej2_entrada", "s1", "ENTRADA", "A3:I11", ["B10:I10"]),
    ("ej2_estado", "s1", "ENTRADA", "Y3:AC11", ["Z10:Z10", "AB10:AC10"]),
    ("ej3_reales", "s2", "ENTRADA", "J3:U11", ["K10:K10", "M10:M10", "S10:S10"]),
    ("ej3_estado", "s2", "ENTRADA", "Y3:AC11", ["Y10:AA10", "AB10:AC10"]),
    ("ej3_inicio", "s2", "INICIO", "B35:C45", ["C37:C37"]),
    ("ej4_precios", "s3", "PRECIOS", ("A5:O8", "A,B,C,D,E,F,O"), ["A7:A7", "E7:E7", "O7:O7"]),
    ("ej4_calculo", "s3", "CALCULO", ("ID contenedor", "Fecha de arribo", "Precios vigentes desde", "Litros al puerto (aplicados)",
                                      "Combustible al puerto", "TOTAL DISTRIBUCIÓN",
                                      "UTILIDAD ANTES DE IMPUESTO SOBRE UTILIDADES"), ["10:11"]),
    ("ej5_comercial_revisar", "s3", "COMERCIAL", "A31:I41", ["G36:G39"]),
    ("ej5_costos", "s4", "COSTOS_FIJOS", "A5:I12", ["H10:H10"]),
    ("ej5_comercial_ok", "s4", "COMERCIAL", "A11:I41", ["G19:G19", "G24:G24", "G31:G31", "G39:G39"]),
    ("ej5_resumen", "s4", "RESUMEN_MES", ("A3:P9", "A,B,C,F,G,H,I,K,M,N,O"), ["K7:K7", "O7:O8"]),
    ("ej6_dashboard", "s5", "DASHBOARD", ("A1:H19", "B,C,G,H"), ["C4:C5", "C8:C18"]),
    ("ej7_simulador", "s6", "SIMULADOR", "B3:F27", ["D11:D11", "E11:E11"]),
    ("ej7_resultado", "s6", "SIMULADOR", "B29:F48", ["E38:E41", "F44:F46"]),
    ("ej8_vacio", "s7", "ENTRADA", "A3:I11", ["A6:I9"]),
    ("ej8_inicio", "s7", "INICIO", "B35:C45", ["C36:C42"]),
]

# valores que la guía cita en el texto: macro -> (estado, hoja, celda | (encabezado, fila), formato)
VALS = {
    "UAItotal": ("s0", "DASHBOARD", "C16", "usd0"),
    "UNtotal": ("s0", "DASHBOARD", "C18", "usd0"),
    "Ingtotal": ("s0", "DASHBOARD", "C10", "usd0"),
    "MCpct": ("s0", "DASHBOARD", "C13", "pct"),
    "CostoKg": ("s0", "DASHBOARD", "H10", "usd3"),
    "IngKg": ("s0", "DASHBOARD", "H9", "usd3"),
    "PEkg": ("s0", "DASHBOARD", "H13", "kg"),
    "MSeg": ("s0", "DASHBOARD", "H15", "pct"),
    "UAIocc": ("s0", "DASHBOARD", "J22", "usd0"),
    "UAIcen": ("s0", "DASHBOARD", "J23", "usd0"),
    "UAIori": ("s0", "DASHBOARD", "J24", "usd0"),
    "UAIcUno": ("s0", "ENTRADA", "AC6", "usd2"),
    "UAIcTres": ("s0", "ENTRADA", "AC8", "usd2"),
    "AlertaTres": ("s0", "ENTRADA", "AB8", "text"),
    "UAIsep": ("s0", "RESUMEN_MES", "O6", "usd0"),
    "UAIoct": ("s0", "RESUMEN_MES", "O7", "usd0"),
    "FijosMes": ("s0", "RESUMEN_MES", "K6", "usd2"),
    "PEsep": ("s0", "RESUMEN_MES", "T6", "kg"),
    "PEcontSep": ("s0", "RESUMEN_MES", "U6", "num1"),
    "SimUAIcont": ("s0", "SIMULADOR", "F38", "usd0"),
    "SimTminOri": ("s0", "SIMULADOR", "E41", "usd3"),
    "SimTminOcc": ("s0", "SIMULADOR", "C41", "usd3"),
    "SimUAImes": ("s0", "SIMULADOR", "F44", "usd0"),
    "SimPE": ("s0", "SIMULADOR", "F48", "num2"),
    "SimGrid": ("s0", "SIMULADOR", "E56", "usd0"),
    "SimGridMas": ("s0", "SIMULADOR", "G56", "usd0"),
    "AlertasIni": ("s0", "INICIO", ("Contenedores con alertas", "C"), "int"),
    "AbiertosIni": ("s0", "INICIO", ("Contenedores abiertos", "C"), "int"),
    "SimUAIori": ("s0", "SIMULADOR", "E38", "usd0"),
    "SimMgOri": ("s0", "SIMULADOR", "E39", "pct"),
    "EjDosUAI": ("s1", "ENTRADA", "AC10", "usd2"),
    "EjDosUAItres": ("s1", "ENTRADA", "AC8", "usd2"),
    "EjTresUAI": ("s2", "ENTRADA", "AC10", "usd2"),
    "EjTresAlerta": ("s2", "ENTRADA", "AB10", "text"),
    "EjTresAlertas": ("s2", "INICIO", ("Contenedores con alertas", "C"), "int"),
    "EjCuatroUAIcinco": ("s3", "ENTRADA", "AC10", "usd2"),
    "EjCuatroUAIseis": ("s3", "ENTRADA", "AC11", "usd2"),
    "EjCuatroCombSeis": ("s3", "CALCULO", ("Combustible al puerto", 11), "usd2"),
    "EjCuatroCombCinco": ("s3", "CALCULO", ("Combustible al puerto", 10), "usd2"),
    "EjCincoDif": ("s3", "COMERCIAL", "G38", "usd2"),
    "EjCincoReal": ("s3", "COMERCIAL", "G37", "usd2"),
    "EjCincoFijosOct": ("s4", "RESUMEN_MES", "K7", "usd2"),
    "EjCincoUAIoct": ("s4", "RESUMEN_MES", "O7", "usd0"),
    "EjCincoUAInov": ("s4", "RESUMEN_MES", "O8", "usd0"),
    "EjCincoPagoTotal": ("s4", "COMERCIAL", "I33", "usd2"),
    "EjCincoPagoTotalAntes": ("s3", "COMERCIAL", "I33", "usd2"),
    "EjCincoDeclarada": ("s3", "COMERCIAL", "G36", "usd2"),
    "EjCincoHabAntes": ("s0", "COMERCIAL", "G19", "plain2"),
    "EjCincoHab": ("s4", "COMERCIAL", "G19", "plain2"),
    "EjCincoVcAntes": ("s0", "COMERCIAL", "G24", "plain2"),
    "EjCincoVc": ("s4", "COMERCIAL", "G24", "plain2"),
    "EjCincoScAntes": ("s0", "COMERCIAL", "G31", "plain2"),
    "EjCincoSc": ("s4", "COMERCIAL", "G31", "plain2"),
    "EjCincoFijosOctAntes": ("s3", "RESUMEN_MES", "K7", "usd2"),
    "EjOchoUAI": ("s7", "INICIO", ("Utilidad antes de impuesto acumulada", "C"), "usd0"),
    "EjOchoConc": ("s7", "INICIO", ("Conciliación de comercialización", "C"), "text"),
    "EjOchoFijos": ("s7", "RESUMEN_MES", "K6", "usd2"),
    "EjSeisCont": ("s5", "DASHBOARD", "C8", "int"),
    "EjSeisUAI": ("s5", "DASHBOARD", "C16", "usd0"),
    "EjSeisMargen": ("s5", "DASHBOARD", "H8", "pct"),
    "EjSieteUAIori": ("s6", "SIMULADOR", "E38", "usd0"),
    "EjSieteMgOri": ("s6", "SIMULADOR", "E39", "pct"),
    "EjSieteUAImes": ("s6", "SIMULADOR", "F44", "usd0"),
}


# ------------------------------------------------------------------ motor
def build_state(name):
    wb = load_workbook(LIBRO)
    for fn in ESTADOS[name]:
        fn(wb)
    path = TMP / f"{name}.xlsx"
    wb.save(path)
    return path


def soffice_pdf(xlsx, outdir, profile):
    subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
                    "--convert-to", "pdf", "--outdir", str(outdir), str(xlsx)],
                   env={**os.environ, "SAL_USE_VCLPLUGIN": "svp"}, capture_output=True, timeout=240, check=False)
    pdf = outdir / (xlsx.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError(f"LibreOffice no generó {pdf}")
    return pdf


def mark(ws, ref, color):
    side = Side(style="thick", color=color)
    if ref.replace(":", "").isdigit():          # filas completas, p. ej. "10:11" (solo para CALCULO)
        a, b = (int(x) for x in ref.split(":"))
        cols = [c for c in range(1, ws.max_column + 1)
                if not ws.column_dimensions[get_column_letter(c)].hidden]
        c1, c2, r1, r2 = cols[0], cols[-1], a, b
    else:
        c1, r1, c2, r2 = range_boundaries(ref)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            if r not in (r1, r2) and c not in (c1, c2):
                continue
            cell = ws.cell(r, c)
            b = copy(cell.border)
            cell.border = Border(left=side if c == c1 else b.left, right=side if c == c2 else b.right,
                                 top=side if r == r1 else b.top, bottom=side if r == r2 else b.bottom)


MARK_COLORS = [f"FF00{i:02X}" for i in range(12, 130, 10)]   # rojos casi idénticos, distinguibles por píxel


def shoot(job):
    name, state, sheet, rng, marks = job
    wb = load_workbook(TMP / f"{state}.xlsx")
    for ws in wb.worksheets:
        if ws.title != sheet:
            ws.sheet_state = "hidden"
    ws = wb[sheet]
    wb.active = wb.sheetnames.index(sheet)
    if isinstance(rng, tuple) and len(rng) == 2 and ":" in rng[0]:   # rango con solo algunas columnas visibles
        area, visibles = rng[0], {c.strip() for c in rng[1].split(",")}
        c1, r1, c2, r2 = range_boundaries(area)
        for c in range(c1, c2 + 1):
            L = get_column_letter(c)
            if L not in visibles:
                ws.column_dimensions[L].hidden = True
        rng = area
    elif isinstance(rng, tuple):                 # columnas elegidas de CALCULO
        keep = {col_of(ws, h) for h in rng}
        last = max(column_index_from_string(c) for c in keep)
        for c in range(1, last + 1):
            L = get_column_letter(c)
            ws.column_dimensions[L].hidden = L not in keep
            if L in keep:
                ws.column_dimensions[L].width = 14 if L != col_of(ws, "ALERTAS") else 40
        ws.row_dimensions[5].height = 70
        first = min(column_index_from_string(c) for c in keep)
        rng = f"A5:{get_column_letter(last)}11"
        marks = [f"{get_column_letter(first)}{m.split(':')[0]}:{get_column_letter(last)}{m.split(':')[1]}" for m in marks]
    ws.print_area = rng
    ws.print_options.headings = True
    ws.print_options.gridLines = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for side in ("left", "right", "top", "bottom"):
        setattr(ws.page_margins, side, 0.2)
    ws.page_margins.header = ws.page_margins.footer = 0
    for i, m in enumerate(marks):
        mark(ws, m, MARK_COLORS[i])
    work = TMP / name
    work.mkdir(exist_ok=True)
    x = work / f"{name}.xlsx"
    wb.save(x)
    pdf = soffice_pdf(x, work, TMP / f"prof_{os.getpid()}")
    subprocess.run(["pdftoppm", "-r", "170", "-png", "-singlefile", "-f", "1", "-l", "1", str(pdf), str(work / "p")], check=True)
    img = Image.open(work / "p.png").convert("RGB")
    img = annotate(img, len(marks))
    img = trim(img)
    img.save(IMG / f"{name}.png", optimize=True)
    return name


def trim(img, pad=12):
    px = img.load()
    w, h = img.size
    xs, ys = [], []
    for y in range(0, h, 2):
        for x in range(0, w, 2):
            r, g, b = px[x, y]
            if r < 245 or g < 245 or b < 245:
                xs.append(x)
                ys.append(y)
    if not xs:
        return img
    box = (max(min(xs) - pad, 0), max(min(ys) - pad, 0), min(max(xs) + pad, w), min(max(ys) + pad, h))
    return img.crop(box)


def annotate(img, n):
    """Dibuja un círculo numerado en la esquina de cada recuadro rojo."""
    if n == 0:
        return img
    px = img.load()
    w, h = img.size
    boxes = {}
    targets = {i: tuple(int(MARK_COLORS[i][k:k + 2], 16) for k in (0, 2, 4)) for i in range(n)}
    for y in range(h):
        for x in range(w):
            p = px[x, y]
            if p[0] < 250 or p[1] > 6 or p[2] > 140:
                continue
            for i, t in targets.items():
                if abs(p[2] - t[2]) <= 2:
                    b = boxes.get(i)
                    boxes[i] = (x, y, x, y) if b is None else (min(b[0], x), min(b[1], y), max(b[2], x), max(b[3], y))
                    break
    # margen para que los círculos no queden cortados
    m = 30
    canvas = Image.new("RGB", (w + 2 * m, h + 2 * m), "white")
    canvas.paste(img, (m, m))
    d = ImageDraw.Draw(canvas)
    font = ImageFont.truetype(FONT, 22)
    for i in range(n):
        if i not in boxes:
            print(f"  aviso: recuadro {i + 1} no localizado")
            continue
        x0, y0 = boxes[i][0] + m, boxes[i][1] + m
        r = 17
        cx, cy = x0 - 4, y0 - 4
        d.ellipse((cx - r, cy - r, cx + r, cy + r), fill=(200, 0, 0), outline="white", width=3)
        d.text((cx, cy), str(i + 1), fill="white", font=font, anchor="mm")
    return canvas


def fmt(v, kind):
    if v is None or v == "":
        return "---" if kind != "text" else "(sin alertas)"
    if kind == "text":
        return str(v).strip().replace("%", r"\%").replace("&", r"\&")
    if isinstance(v, str):
        return v
    neg = v < 0
    a = abs(v)
    s = {
        "usd0": lambda: f"\\${a:,.0f}", "usd2": lambda: f"\\${a:,.2f}", "usd3": lambda: f"\\${a:,.3f}",
        "pct": lambda: f"{a * 100:.1f}\\%", "kg": lambda: f"{a:,.0f}", "int": lambda: f"{a:.0f}",
        "num1": lambda: f"{a:,.1f}", "num2": lambda: f"{a:,.2f}", "plain2": lambda: f"{a:.2f}",
    }[kind]()
    return f"({s})" if neg and kind.startswith("usd") else ("-" + s if neg else s)


def main():
    if not RECALC:
        sys.exit("Defina RECALC con la ruta de recalc.py")
    IMG.mkdir(parents=True, exist_ok=True)
    datos = {}
    for st in ESTADOS:
        p = build_state(st)
        q = TMP / f"{st}_val.xlsx"
        shutil.copy(p, q)
        out = subprocess.run([sys.executable, RECALC, str(q), "300"], capture_output=True, text=True)
        res = json.loads(out.stdout)
        if res.get("total_errors", 1) != 0:
            sys.exit(f"Estado {st}: errores de fórmula {res}")
        datos[st] = load_workbook(q, data_only=True)
        print(f"estado {st}: {res['total_formulas']} fórmulas, 0 errores")
    lines = ["% Archivo generado por build/capturas.py. No editar a mano.", ""]
    for macro, (st, sh, ref, kind) in VALS.items():
        ws = datos[st][sh]
        if isinstance(ref, tuple):
            a, b = ref
            if isinstance(b, int):
                ref = f"{col_of(ws, a)}{b}"
            else:
                ref = f"{b}{find_row(ws, 'B', a)}"
        lines.append(f"\\newcommand{{\\v{macro}}}{{{fmt(ws[ref].value, kind)}}}")
    VALORES.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"{len(VALS)} valores escritos en {VALORES}")
    with cf.ProcessPoolExecutor(max_workers=4) as ex:
        for n in ex.map(shoot, CAPTURAS):
            print("  captura", n)
    shutil.rmtree(TMP, ignore_errors=True)


if __name__ == "__main__":
    main()
