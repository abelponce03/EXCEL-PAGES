# -*- coding: utf-8 -*-
"""Genera el libro Analisis_Rentabilidad_Paqueteria.xlsx.

Uso:  python build/generar_excel.py [ruta_salida]
"""
import datetime as dt
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from contenido import ANALISIS, GUIA  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
ARGS = [a for a in sys.argv[1:] if not a.startswith("--")]
OUT = Path(ARGS[0]) if ARGS else ROOT / "Analisis_Rentabilidad_Paqueteria.xlsx"
SIN_EJEMPLOS = "--vacio" in sys.argv

# ---------------------------------------------------------------- estilos
FONT = "Arial"
NAVY = "1F3864"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_H2 = Font(name=FONT, size=12, bold=True, color=NAVY)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_GRP = PatternFill("solid", fgColor="D9E1F2")
FILL_INPUT = PatternFill("solid", fgColor="FFF9DB")
FILL_TOTAL = PatternFill("solid", fgColor="F2F2F2")
FILL_KPI = PatternFill("solid", fgColor="EEF3FA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

FMT = {
    "usd": '$#,##0.00;[Red]($#,##0.00);"-"',
    "usd0": '$#,##0;[Red]($#,##0);"-"',
    "usdkg": '$0.000;[Red]($0.000);"-"',
    "pct": '0.0%;[Red]-0.0%;"-"',
    "kg": '#,##0;[Red]-#,##0;"-"',
    "num": '#,##0.0;[Red]-#,##0.0;"-"',
    "int": '0',
    "date": 'dd/mm/yyyy;;"-"',
    "mon": 'mmm-yyyy',
    "text": '@',
}

DATA0, DATAN = 6, 305          # filas de contenedores (300)
PR0, PRN = 6, 55               # filas de vigencias de precios (50)
NMES = 24                      # meses del período

wb = Workbook()
wb.remove(wb.active)


def sheet(name, tab=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    if tab:
        ws.sheet_properties.tabColor = tab
    return ws


def put(ws, ref, value, font=F_BASE, fmt=None, fill=None, align=None, border=False):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = FMT.get(fmt, fmt)
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = BORDER
    return c


def inp(ws, ref, value, fmt=None):
    """Celda de entrada: letra azul, fondo amarillo claro, desbloqueada."""
    c = put(ws, ref, value, F_INPUT, fmt, FILL_INPUT, border=True)
    c.protection = Protection(locked=False)
    return c


def hdr(ws, ref, text, fill=FILL_HDR, font=F_HDR):
    return put(ws, ref, text, font, None, fill, CENTER, True)


def title(ws, text, sub=None):
    put(ws, "A1", text, F_TITLE)
    if sub:
        put(ws, "A2", sub, F_SUB)
    ws.row_dimensions[1].height = 24


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def protect(ws):
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.autoFilter = False


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


# ================================================================ LISTAS
ls = sheet("LISTAS", "7F7F7F")
title(ls, "Listas de validación", "Hoja técnica. Los valores alimentan los desplegables del libro.")
TERRITORIOS = [
    ("Pinar del Río", "Occidente"), ("Artemisa", "Occidente"), ("La Habana", "Occidente"),
    ("Mayabeque", "Occidente"), ("Matanzas", "Occidente"), ("Isla de la Juventud", "Occidente"),
    ("Cienfuegos", "Centro"), ("Villa Clara", "Centro"), ("Sancti Spíritus", "Centro"),
    ("Ciego de Ávila", "Centro"), ("Camagüey", "Centro"),
    ("Las Tunas", "Oriente"), ("Holguín", "Oriente"), ("Granma", "Oriente"),
    ("Santiago de Cuba", "Oriente"), ("Guantánamo", "Oriente"),
    ("Oficina central", "General"),
]
LISTAS = {
    "L_Transitaria": ["ENCI", "A.V", "Otra"],
    "L_Region": ["Occidente", "Centro", "Oriente"],
    "L_SiNo": ["Sí", "No"],
    "L_Prorrateo": ["KG", "CONTENEDOR"],
    "L_BaseCom": ["% DEL INGRESO", "USD POR KG"],
    "L_Estado": ["Abierto", "Cerrado"],
    "L_TipoCont": ["20'", "40'", "40' HC", "Otro"],
    "L_Grupo": ["Grupo central", "Provincial"],
    "L_Centro": ["GENERAL", "OCC", "CEN", "ORI"],
    "L_Elemento": ["Materias primas y materiales", "Combustibles y lubricantes", "Energía",
                   "Gastos de fuerza de trabajo", "Depreciación y amortización", "Otros gastos monetarios"],
    "L_TipoActivo": ["Vehículo propio", "Equipo", "Otro"],
}
col = 1
for nm, vals in LISTAS.items():
    L = CL(col)
    hdr(ls, f"{L}4", nm.replace("L_", ""))
    for i, v in enumerate(vals):
        put(ls, f"{L}{5 + i}", v, border=True)
    name(nm, f"LISTAS!${L}$5:${L}${4 + len(vals)}")
    ls.column_dimensions[L].width = 22
    col += 1
# territorios -> región
TL, RL = CL(col), CL(col + 1)
hdr(ls, f"{TL}4", "Territorio")
hdr(ls, f"{RL}4", "Región")
for i, (t, r) in enumerate(TERRITORIOS):
    put(ls, f"{TL}{5 + i}", t, border=True)
    put(ls, f"{RL}{5 + i}", r, border=True)
name("L_Territorio", f"LISTAS!${TL}$5:${TL}${4 + len(TERRITORIOS)}")
name("L_TerrRegion", f"LISTAS!${RL}$5:${RL}${4 + len(TERRITORIOS)}")
ls.column_dimensions[TL].width = 22
ls.column_dimensions[RL].width = 14


def dv_list(ws, rng, listname, allow_blank=True):
    dv = DataValidation(type="list", formula1=f"={listname}", allow_blank=allow_blank,
                        showErrorMessage=True, errorTitle="Valor no válido",
                        error="Seleccione un valor de la lista.")
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_num(ws, rng, minimo=0):
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=str(minimo),
                        allow_blank=True, showErrorMessage=True, errorTitle="Número no válido",
                        error=f"Escriba un número mayor o igual que {minimo}, sin texto.")
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_date(ws, rng):
    dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)",
                        formula2="DATE(2040,12,31)", allow_blank=True, showErrorMessage=True,
                        errorTitle="Fecha no válida", error="Escriba una fecha dd/mm/aaaa.")
    ws.add_data_validation(dv)
    dv.add(rng)


# ================================================================ PARAMETROS
pa = sheet("PARAMETROS", "C00000")
title(pa, "Parámetros generales, normas técnicas y tributos",
      "Valores estables del modelo. Los precios que cambian con el tiempo van en la hoja PRECIOS (con fecha de vigencia).")
for c, t in zip("ABCDE", ["Nombre (referencia)", "Concepto", "Valor", "Unidad", "Fuente / nota"]):
    hdr(pa, f"{c}4", t)
ILU = "ILUSTRATIVO: sustituir por el dato real."
PARAMS = [
    ("A. GENERALES",),
    ("Empresa", "Nombre de la empresa", "Empresa de Paquetería (editar)", "texto", "Aparece en los encabezados.", None),
    ("Fecha_Inicio", "Primer mes del período de análisis", dt.date(2026, 9, 1), "fecha",
     "Primer mes que se registra. COSTOS_FIJOS y RESUMEN_MES abarcan 24 meses desde esta fecha.", "date"),
    ("Metodo_Prorrateo", "Base de reparto de los costos fijos entre contenedores", "KG", "lista",
     "KG = proporcional al peso (recomendado: un contenedor más cargado absorbe más estructura). CONTENEDOR = partes iguales.", None),
    ("Base_Comision", "Base del pago variable de comercialización", "% DEL INGRESO", "lista",
     "% DEL INGRESO: tasa sobre lo cobrado. USD POR KG: importe por kg entregado. Las tasas están en PRECIOS.", None),
    ("Tolerancia_Comb", "Tolerancia del desvío de combustible frente a la norma", 0.10, "%",
     "Si el consumo real supera la norma en más de este %, se genera una alerta.", "pct"),
    ("Mes_Corte", "Último mes incluido en los resultados (opcional)", None, "fecha",
     "Vacío = último mes con contenedores registrados. Los meses posteriores no cargan costos fijos en RESUMEN_MES ni en DASHBOARD.", "date"),
    ("B. NORMAS TÉCNICAS (se usan cuando no hay dato real)",),
    ("Std_Viajes_Puerto", "Viajes al puerto por contenedor", 1, "viajes", "Un contenedor = un viaje del camión propio.", "num"),
    ("Litros_Puerto_Viaje", "Combustible por viaje al puerto", 60, "L / viaje", "Dato informado por la empresa.", "num"),
    ("Km_Puerto_Viaje", "Km por viaje al puerto (ida y vuelta)", 180, "km / viaje", ILU + " Medir con la hoja de ruta.", "num"),
    ("Rend_Desagrupe", "Rendimiento del desagrupe en la transitaria", 1000, "kg / jornal",
     ILU + " Jornales = kg ÷ rendimiento (redondeado hacia arriba).", "kg"),
    ("Rend_Desagrupe_CT", "Rendimiento del desagrupe en los centros transitorios (Centro y Oriente)", 800, "kg / jornal", ILU, "kg"),
    ("Cap_Camion_OCC", "Capacidad útil del camión de distribución – Occidente", 5000, "kg / viaje", ILU + " Viajes = kg ÷ capacidad (redondeado hacia arriba).", "kg"),
    ("Cap_Camion_CEN", "Capacidad útil del camión de traslado – Centro", 5000, "kg / viaje", ILU, "kg"),
    ("Cap_Camion_ORI", "Capacidad útil del camión de traslado – Oriente", 5000, "kg / viaje", ILU, "kg"),
    ("L_Dist_OCC", "Combustible de distribución en Occidente", 120, "L / viaje", ILU, "num"),
    ("L_Tras_CEN", "Combustible de traslado al centro transitorio – Centro", 250, "L / viaje", ILU, "num"),
    ("L_Dist_CEN", "Combustible de distribución en Centro", 150, "L / viaje", ILU, "num"),
    ("L_Tras_ORI", "Combustible de traslado al centro transitorio – Oriente", 450, "L / viaje", ILU, "num"),
    ("L_Dist_ORI", "Combustible de distribución en Oriente", 180, "L / viaje", ILU, "num"),
    ("C. OTROS COSTOS VARIABLES (sobre el ingreso)",),
    ("Pct_Reclamaciones", "Provisión para pérdidas, daños y reclamaciones", 0.005, "% ingreso",
     "Recomendado. " + ILU + " Use 0 si no aplica.", "pct"),
    ("Pct_Banco", "Comisiones bancarias y costo de cobro", 0.01, "% ingreso",
     "Recomendado. " + ILU + " Use 0 si no aplica.", "pct"),
    ("D. TRIBUTOS (régimen general MIPYME – confirmar con el contador / ONAT)",),
    ("Imp_Ventas_Serv", "Impuesto sobre las ventas y los servicios", 0.10, "% ingreso",
     "Ley 113 del Sistema Tributario; vigente en 2026 según la Ley del Presupuesto.", "pct"),
    ("Contrib_Territorial", "Contribución territorial para el desarrollo local", 0.01, "% ingreso", "Ley 113; vigente en 2026.", "pct"),
    ("Contrib_SS", "Contribución a la seguridad social (empleador)", 0.14, "% salarios", "Solo sobre salarios de trabajadores propios.", "pct"),
    ("Imp_Fuerza_Trabajo", "Impuesto por la utilización de la fuerza de trabajo", 0.05, "% salarios",
     "Solo sobre salarios de trabajadores propios. Confirmar si la empresa tiene alguna bonificación.", "pct"),
    ("Imp_Utilidades", "Impuesto sobre las utilidades", 0.35, "% utilidad",
     "Liquidación anual con pagos a cuenta. En el libro es una estimación.", "pct"),
]
r = 5
for p in PARAMS:
    if len(p) == 1:
        for c in "ABCDE":
            pa[f"{c}{r}"].fill = FILL_GRP
        put(pa, f"A{r}", p[0], F_BOLD, fill=FILL_GRP)
        r += 1
        continue
    nm, desc, val, unit, note, fmt = p
    put(pa, f"A{r}", nm, Font(name=FONT, size=9, color="7F7F7F"), border=True)
    put(pa, f"B{r}", desc, border=True, align=WRAP)
    inp(pa, f"C{r}", val, fmt)
    put(pa, f"D{r}", unit, border=True)
    put(pa, f"E{r}", note, F_SUB, border=True, align=WRAP)
    name(nm, f"PARAMETROS!$C${r}")
    if nm == "Metodo_Prorrateo":
        dv_list(pa, f"C{r}", "L_Prorrateo", False)
    elif nm == "Base_Comision":
        dv_list(pa, f"C{r}", "L_BaseCom", False)
    elif nm in ("Fecha_Inicio", "Mes_Corte"):
        dv_date(pa, f"C{r}")
    elif fmt in ("num", "kg", "pct"):
        dv_num(pa, f"C{r}")
    r += 1
widths(pa, {"A": 22, "B": 58, "C": 24, "D": 13, "E": 70})
pa.freeze_panes = "A5"
protect(pa)

# ================================================================ PRECIOS
pr = sheet("PRECIOS", "C00000")
title(pr, "Precios y tarifas con fecha de vigencia",
      "Cuando cambie un precio, AGREGUE una fila nueva con la fecha desde la que rige (en orden ascendente). "
      "No sobrescriba las anteriores: cada contenedor usa los precios vigentes en su fecha de arribo.")
PRCOLS = [
    ("A", "Vigente desde", "date", 12),
    ("B", "Tarifa Occidente (USD/kg)", "usdkg", 12),
    ("C", "Tarifa Centro (USD/kg)", "usdkg", 12),
    ("D", "Tarifa Oriente (USD/kg)", "usdkg", 12),
    ("E", "Precio diésel (USD/L)", "usdkg", 11),
    ("F", "Cita del puerto (USD/viaje)", "usd", 11),
    ("G", "Dieta del chofer (USD/viaje)", "usd", 11),
    ("H", "Pago por jornal – transitaria (USD)", "usd", 12),
    ("I", "Pago por jornal – centro transitorio (USD)", "usd", 13),
    ("J", "Servicio ENCI (USD/contenedor)", "usd", 12),
    ("K", "Servicio A.V (USD/contenedor)", "usd", 12),
    ("L", "Comisión variable (% del ingreso)", "pct", 12),
    ("M", "Comisión variable (USD/kg)", "usdkg", 12),
    ("N", "Desgaste del vehículo propio (USD/km)", "usdkg", 12),
    ("O", "Documento de soporte / nota", "text", 60),
]
for c, t, f, w in PRCOLS:
    hdr(pr, f"{c}5", t)
    pr.column_dimensions[c].width = w
pr.row_dimensions[5].height = 54
EJ_PR = [dt.date(2026, 1, 1), 0.80, 0.80, 0.80, 2.00, 10, 10, 15, 15, 0, 0, 0.03, 0.03, 0.12,
         "EJEMPLO. Cita 10 USD = dato de la empresa. Diésel 2,00 USD/L = referencia de CUPET 2026. "
         "Resto ILUSTRATIVO: sustituir."]
for i in range(PR0, PRN + 1):
    for j, (c, t, f, w) in enumerate(PRCOLS):
        v = EJ_PR[j] if i == PR0 else None
        inp(pr, f"{c}{i}", v, f)
dv_date(pr, f"A{PR0}:A{PRN}")
dv_num(pr, f"B{PR0}:N{PRN}")
pr.freeze_panes = "B6"
protect(pr)

# ================================================================ ACTIVOS
ac = sheet("ACTIVOS", "C00000")
title(ac, "Activos fijos: depreciación, mantenimiento y seguros",
      "Método lineal. La depreciación se calcula desde el mes de alta hasta el fin de la vida útil o la fecha de baja. "
      "Confirme las tasas con la norma contable vigente.")
ACOLS = [("A", "Nº", 5), ("B", "Descripción / chapa", 30), ("C", "Tipo", 15), ("D", "Fecha de alta", 12),
         ("E", "Valor de adquisición (USD)", 14), ("F", "Valor residual (USD)", 12), ("G", "Vida útil (años)", 9),
         ("H", "Depreciación anual (USD)", 13), ("I", "Depreciación mensual (USD)", 13),
         ("J", "Fin de la depreciación", 12), ("K", "Fecha de baja (si se vendió o retiró)", 13),
         ("L", "Mantenimiento preventivo (USD/mes)", 14), ("M", "Seguros, licencias e inspecciones (USD/año)", 15),
         ("N", "Seguros y licencias (USD/mes)", 13), ("O", "Observaciones", 40)]
for c, t, w in ACOLS:
    hdr(ac, f"{c}5", t)
    ac.column_dimensions[c].width = w
ac.row_dimensions[5].height = 54
EJ_AC = {
    6: ["Camión porta-contenedor (recogida en el puerto) – EJEMPLO", "Vehículo propio", dt.date(2025, 1, 1),
        45000, 5000, 8, None, 150, 1020, "ILUSTRATIVO: sustituir por el valor contable real."],
    7: ["Computadoras y equipos de oficina – EJEMPLO", "Equipo", dt.date(2026, 1, 1), 1500, 0, 5, None, 0, 0, "ILUSTRATIVO"],
}
A0, AN = 6, 15
for i in range(A0, AN + 1):
    put(ac, f"A{i}", i - 5, border=True, align=CENTER)
    ej = EJ_AC.get(i, [None] * 10)
    inp(ac, f"B{i}", ej[0])
    inp(ac, f"C{i}", ej[1])
    inp(ac, f"D{i}", ej[2], "date")
    inp(ac, f"E{i}", ej[3], "usd")
    inp(ac, f"F{i}", ej[4], "usd")
    inp(ac, f"G{i}", ej[5], "num")
    put(ac, f"H{i}", f'=IF(OR(D{i}="",N(G{i})=0),0,(E{i}-F{i})/G{i})', fmt="usd", border=True)
    put(ac, f"I{i}", f"=H{i}/12", fmt="usd", border=True)
    put(ac, f"J{i}", f'=IF(OR(D{i}="",N(G{i})=0),0,EDATE(D{i},G{i}*12))', fmt="date", border=True)
    inp(ac, f"K{i}", ej[6], "date")
    inp(ac, f"L{i}", ej[7], "usd")
    inp(ac, f"M{i}", ej[8], "usd")
    put(ac, f"N{i}", f"=N(M{i})/12", fmt="usd", border=True)
    inp(ac, f"O{i}", ej[9])
T = AN + 1
put(ac, f"B{T}", "TOTAL", F_BOLD, fill=FILL_TOTAL, border=True)
for c in "EFHILN":
    put(ac, f"{c}{T}", f"=SUM({c}{A0}:{c}{AN})", F_BOLD, "usd", FILL_TOTAL, border=True)
dv_list(ac, f"C{A0}:C{AN}", "L_TipoActivo")
dv_date(ac, f"D{A0}:D{AN}")
dv_date(ac, f"K{A0}:K{AN}")
dv_num(ac, f"E{A0}:G{AN}")
dv_num(ac, f"L{A0}:M{AN}")
ac.freeze_panes = "C6"
protect(ac)
AC_RNG = lambda c: f"ACTIVOS!${c}${A0}:${c}${AN}"  # noqa: E731

# ================================================================ COMERCIAL (roster; liquidación al final)
co = sheet("COMERCIAL", "C00000")
CO0, CON = 12, 32   # 21 gestores
CO_TOT = CON + 1

# ================================================================ COSTOS_FIJOS
cf = sheet("COSTOS_FIJOS", "C00000")
title(cf, "Costos fijos mensuales por centro de costo",
      "La columna «Valor base» se copia a cada mes. Si el gasto real de un mes fue distinto, escriba el valor real "
      "en la celda de ese mes. Las filas en negro se calculan solas.")
MC0 = 7                        # columna G = primer mes
MCN = MC0 + NMES - 1           # columna AD
ML0, MLN = CL(MC0), CL(MCN)
for c, t, w in [("A", "Código", 7), ("B", "Concepto", 46), ("C", "Centro de costo", 10),
                ("D", "Elemento de gasto", 24), ("E", "¿Salario de trabajador propio?", 11),
                ("F", "Valor base mensual (USD)", 13)]:
    hdr(cf, f"{c}5", t)
    cf.column_dimensions[c].width = w
for k in range(NMES):
    L = CL(MC0 + k)
    c = put(cf, f"{L}5", f"=EDATE(Fecha_Inicio,{k})", F_HDR, "mon", FILL_HDR, CENTER, True)
    cf.column_dimensions[L].width = 11
TCOL = CL(MCN + 1)
hdr(cf, f"{TCOL}5", "Total período")
cf.column_dimensions[TCOL].width = 13
cf.row_dimensions[5].height = 42

GTF = "Gastos de fuerza de trabajo"
OGM = "Otros gastos monetarios"
CONCEPTOS = [
    ("CF01", "Salario del chofer del vehículo propio (puerto)", "GENERAL", GTF, "Sí", 60),
    ("CF02", "Salarios de dirección, economía y contabilidad", "GENERAL", GTF, "Sí", 250),
    ("CF03", "Comercialización – parte fija (21 gestores, desde COMERCIAL)", "GENERAL", GTF, "Sí", f"=COMERCIAL!F{CO_TOT}"),
    ("CF04", "Alquiler de local / almacén", "GENERAL", OGM, "No", 200),
    ("CF05", "Electricidad y agua", "GENERAL", "Energía", "No", 40),
    ("CF06", "Comunicaciones (teléfono, internet)", "GENERAL", OGM, "No", 30),
    ("CF07", "Materiales de oficina y limpieza", "GENERAL", "Materias primas y materiales", "No", 15),
    ("CF08", "Servicios contables, legales y otros externos", "GENERAL", OGM, "No", 50),
    ("CF09", "Gastos bancarios fijos (mantenimiento de cuenta)", "GENERAL", OGM, "No", 10),
    ("CF10", "Depreciación de activos fijos (desde ACTIVOS)", "GENERAL", "Depreciación y amortización", "No", "DEP"),
    ("CF11", "Mantenimiento preventivo de la flota propia (desde ACTIVOS)", "GENERAL", "Materias primas y materiales", "No", "MANT"),
    ("CF12", "Seguros, licencias e inspecciones de la flota (desde ACTIVOS)", "GENERAL", OGM, "No", "SEG"),
    ("CF13", "Otros gastos fijos generales", "GENERAL", OGM, "No", 0),
    ("CF14", "Occidente – chofer (servicio contratado)", "OCC", OGM, "No", 120),
    ("CF15", "Occidente – ayudante (servicio contratado)", "OCC", OGM, "No", 80),
    ("CF16", "Occidente – otros gastos fijos", "OCC", OGM, "No", 0),
    ("CF17", "Centro – chofer (servicio contratado)", "CEN", OGM, "No", 120),
    ("CF18", "Centro – ayudante (servicio contratado)", "CEN", OGM, "No", 80),
    ("CF19", "Centro – otros gastos fijos (centro transitorio)", "CEN", OGM, "No", 0),
    ("CF20", "Oriente – chofer (servicio contratado)", "ORI", OGM, "No", 120),
    ("CF21", "Oriente – ayudante (servicio contratado)", "ORI", OGM, "No", 80),
    ("CF22", "Oriente – otros gastos fijos (centro transitorio)", "ORI", OGM, "No", 0),
]
CF0 = 6
CFN = CF0 + len(CONCEPTOS) - 1
ACT_SUM = {
    "DEP": ("I", True), "MANT": ("L", False), "SEG": ("N", False),
}
for i, (code, desc, pool, elem, sal, base) in enumerate(CONCEPTOS):
    rr = CF0 + i
    put(cf, f"A{rr}", code, border=True)
    if base in ACT_SUM or code == "CF03":
        put(cf, f"B{rr}", desc, border=True)
    else:
        inp(cf, f"B{rr}", desc)
    inp(cf, f"C{rr}", pool)
    inp(cf, f"D{rr}", elem)
    inp(cf, f"E{rr}", sal)
    if base in ACT_SUM:
        acol, is_dep = ACT_SUM[base]
        put(cf, f"F{rr}", f"=ACTIVOS!{acol}{T}", F_LINK, "usd", border=True)
        for k in range(NMES):
            L = CL(MC0 + k)
            m = f"{L}$5"
            cond = (f"({AC_RNG('D')}>0)*({AC_RNG('D')}<=EOMONTH({m},0))"
                    f"*((({AC_RNG('K')}=0)+({AC_RNG('K')}>{m}))>0)")
            if is_dep:
                cond += f"*({AC_RNG('J')}>{m})"
            put(cf, f"{L}{rr}", f"=SUMPRODUCT({cond}*{AC_RNG(acol)})", fmt="usd", border=True)
    else:
        if isinstance(base, str):
            put(cf, f"F{rr}", base, F_LINK, "usd", border=True)
        else:
            inp(cf, f"F{rr}", base, "usd")
        for k in range(NMES):
            L = CL(MC0 + k)
            inp(cf, f"{L}{rr}", f"=$F{rr}", "usd")
    put(cf, f"{TCOL}{rr}", f"=SUM({ML0}{rr}:{MLN}{rr})", F_BOLD, "usd", border=True)
dv_list(cf, f"C{CF0}:C{CFN}", "L_Centro", False)
dv_list(cf, f"D{CF0}:D{CFN}", "L_Elemento")
dv_list(cf, f"E{CF0}:E{CFN}", "L_SiNo", False)

# tributos sobre salarios por centro de costo
TX0 = CFN + 2
put(cf, f"B{TX0 - 1}", "Tributos sobre salarios de trabajadores propios (calculados)", F_BOLD)
row = TX0
for pool in ["GENERAL", "OCC", "CEN", "ORI"]:
    for code, desc, rate in [("SS", "Contribución a la seguridad social", "Contrib_SS"),
                             ("FT", "Impuesto por utilización de la fuerza de trabajo", "Imp_Fuerza_Trabajo")]:
        put(cf, f"A{row}", f"{code}-{pool}", border=True)
        put(cf, f"B{row}", f"{desc} – {pool}", border=True)
        put(cf, f"C{row}", pool, border=True)
        put(cf, f"D{row}", GTF, border=True)
        put(cf, f"E{row}", "—", border=True, align=CENTER)
        put(cf, f"F{row}", f'={rate}*SUMIFS($F${CF0}:$F${CFN},$C${CF0}:$C${CFN},$C{row},$E${CF0}:$E${CFN},"Sí")',
            fmt="usd", border=True)
        for k in range(NMES):
            L = CL(MC0 + k)
            put(cf, f"{L}{row}",
                f'={rate}*SUMIFS({L}${CF0}:{L}${CFN},$C${CF0}:$C${CFN},$C{row},$E${CF0}:$E${CFN},"Sí")',
                fmt="usd", border=True)
        put(cf, f"{TCOL}{row}", f"=SUM({ML0}{row}:{MLN}{row})", F_BOLD, "usd", border=True)
        row += 1
TXN = row - 1
TOT0 = TXN + 2
CF_TOT = {}
for j, (pool, label) in enumerate([("GENERAL", "TOTAL FIJOS GENERALES (se reparten entre todos los kg)"),
                                   ("OCC", "TOTAL FIJOS OCCIDENTE"), ("CEN", "TOTAL FIJOS CENTRO"),
                                   ("ORI", "TOTAL FIJOS ORIENTE"), (None, "TOTAL COSTOS FIJOS DEL MES")]):
    rr = TOT0 + j
    CF_TOT[pool or "ALL"] = rr
    put(cf, f"B{rr}", label, F_BOLD, fill=FILL_TOTAL, border=True)
    for c in "ACDE":
        put(cf, f"{c}{rr}", pool if c == "C" else None, F_BOLD, fill=FILL_TOTAL, border=True)
    for k in list(range(NMES)) + [NMES]:
        L = CL(MC0 + k) if k < NMES else TCOL
        if pool:
            f = f'=SUMIF($C${CF0}:$C${TXN},"{pool}",{L}${CF0}:{L}${TXN})'
        else:
            f = f"=SUM({L}{TOT0}:{L}{TOT0 + 3})"
        put(cf, f"{L}{rr}", f, F_BOLD, "usd", FILL_TOTAL, border=True)
    if pool:
        f = f'=SUMIF($C${CF0}:$C${TXN},"{pool}",$F${CF0}:$F${TXN})'
    else:
        f = f"=SUM(F{TOT0}:F{TOT0 + 3})"
    put(cf, f"F{rr}", f, F_BOLD, "usd", FILL_TOTAL, border=True)
cf.freeze_panes = "G6"
protect(cf)
CF_MONTHS = f"COSTOS_FIJOS!${ML0}$5:${MLN}$5"


def cf_row(pool):
    rr = CF_TOT[pool]
    return f"COSTOS_FIJOS!${ML0}${rr}:${MLN}${rr}"


# ================================================================ ENTRADA
en = sheet("ENTRADA", "2E75B6")
title(en, "Registro de contenedores (ENTRADA DE DATOS)",
      "Una fila por contenedor. Columnas «(opcional)»: si se dejan vacías se usa la norma técnica de PARAMETROS. "
      "Para borrar datos use la tecla Supr; nunca elimine filas.")
ECOLS = [
    # letra, encabezado, grupo, formato, ancho, validación, comentario
    ("A", "Nº", "IDENTIFICACIÓN", "int", 5, None, None),
    ("B", "ID contenedor", "IDENTIFICACIÓN", "text", 18, None, "Número del contenedor o del BL. Obligatorio: sin ID la fila no se calcula."),
    ("C", "Fecha de arribo", "IDENTIFICACIÓN", "date", 11, "date", "Determina el mes y los precios vigentes."),
    ("D", "Transitaria", "IDENTIFICACIÓN", None, 10, "L_Transitaria", "ENCI o A.V"),
    ("E", "Tipo", "IDENTIFICACIÓN", None, 8, "L_TipoCont", None),
    ("F", "Nº de bultos", "IDENTIFICACIÓN", "kg", 9, "num", "Informativo."),
    ("G", "kg Occidente", "VOLUMEN (kg)", "kg", 11, "num", "kg destinados a Pinar del Río, Artemisa, La Habana, Mayabeque, Matanzas e Isla de la Juventud."),
    ("H", "kg Centro", "VOLUMEN (kg)", "kg", 11, "num", "kg destinados a Cienfuegos, Villa Clara, Sancti Spíritus, Ciego de Ávila y Camagüey."),
    ("I", "kg Oriente", "VOLUMEN (kg)", "kg", 11, "num", "kg destinados a Las Tunas, Holguín, Granma, Santiago de Cuba y Guantánamo."),
    ("J", "Viajes al puerto (opcional)", "ETAPA 1 · PUERTO", "num", 9, "num", "Vacío = norma (1 viaje)."),
    ("K", "Litros reales al puerto (opcional)", "ETAPA 1 · PUERTO", "num", 10, "num", "Total de litros del contenedor según los vales. Vacío = viajes × 60 L."),
    ("L", "Dieta pagada (USD, opcional)", "ETAPA 1 · PUERTO", "usd", 10, "num", "Vacío = viajes × dieta vigente."),
    ("M", "Jornales del desagrupe (opcional)", "ETAPA 2 · TRANSITARIA", "num", 10, "num",
     "Jornales = trabajadores × días. Vacío = kg ÷ rendimiento."),
    ("N", "Estadía / almacenaje / demora (USD)", "ETAPA 2 · TRANSITARIA", "usd", 11, "num", "Cargos por exceder los días libres."),
    ("O", "Litros de distribución en Occidente (opcional)", "ETAPA 3 · DISTRIBUCIÓN", "num", 11, "num", None),
    ("P", "Litros de traslado a Centro (opcional)", "ETAPA 3 · DISTRIBUCIÓN", "num", 11, "num", None),
    ("Q", "Litros de distribución en Centro (opcional)", "ETAPA 3 · DISTRIBUCIÓN", "num", 11, "num", None),
    ("R", "Jornales del centro transitorio en Centro (opcional)", "ETAPA 3 · DISTRIBUCIÓN", "num", 11, "num", None),
    ("S", "Litros de traslado a Oriente (opcional)", "ETAPA 3 · DISTRIBUCIÓN", "num", 11, "num", None),
    ("T", "Litros de distribución en Oriente (opcional)", "ETAPA 3 · DISTRIBUCIÓN", "num", 11, "num", None),
    ("U", "Jornales del centro transitorio en Oriente (opcional)", "ETAPA 3 · DISTRIBUCIÓN", "num", 11, "num", None),
    ("V", "Otros gastos directos (USD)", "OTROS", "usd", 11, "num", "Gastos del contenedor que no tienen columna propia."),
    ("W", "Detalle de otros gastos", "OTROS", "text", 22, None, None),
    ("X", "Otros ingresos (USD)", "OTROS", "usd", 11, "num", "Cargos extra cobrados: entrega a domicilio, sobrepeso, etc."),
    ("Y", "Fecha de fin de la distribución", "CONTROL", "date", 11, "date", "Para medir los días de ciclo."),
    ("Z", "Estado", "CONTROL", None, 9, "L_Estado", "Abierto mientras falten datos reales; Cerrado al completar."),
    ("AA", "Observaciones / documentos de soporte", "CONTROL", "text", 30, None, None),
    ("AB", "Alertas (automático)", "RESULTADO", "text", 38, None, None),
    ("AC", "Utilidad antes de impuesto (automático)", "RESULTADO", "usd", 14, None, None),
]
# encabezados de grupo
groups = {}
for c, h, g, *_ in ECOLS:
    groups.setdefault(g, []).append(c)
for g, cols in groups.items():
    a, b = cols[0], cols[-1]
    if a != b:
        en.merge_cells(f"{a}4:{b}4")
    hdr(en, f"{a}4", g, FILL_GRP, F_BOLD)
for c, h, g, f, w, v, com in ECOLS:
    hdr(en, f"{c}5", h)
    en.column_dimensions[c].width = w
    if com:
        en[f"{c}5"].comment = Comment(com, "Modelo")
en.row_dimensions[5].height = 66

EJ_EN = [
    ["CONT-EJEMPLO-01", dt.date(2026, 9, 8), "ENCI", "40' HC", 1050, 5200, 2600, 2300, None, 62, None, 11, 0,
     None, None, None, None, None, None, None, 0, None, 0, dt.date(2026, 9, 18), "Cerrado", "EJEMPLO – borrar"],
    ["CONT-EJEMPLO-02", dt.date(2026, 9, 22), "A.V", "40'", 900, 4600, 2100, 2700, None, None, None, None, 120,
     None, None, None, None, 480, None, None, 0, None, 0, dt.date(2026, 10, 3), "Cerrado", "EJEMPLO – 2 días de estadía"],
    ["CONT-EJEMPLO-03", dt.date(2026, 10, 6), "ENCI", "40' HC", 1180, 6100, 2900, 2400, None, None, None, None, 0,
     150, None, None, None, None, 260, None, 35, "Reempaque de bultos dañados", 150, dt.date(2026, 10, 16), "Cerrado",
     "EJEMPLO – desvío de combustible en Oriente"],
    ["CONT-EJEMPLO-04", dt.date(2026, 10, 20), "A.V", "40'", 860, 3800, 1500, 3900, None, None, None, None, 0,
     None, None, None, None, None, None, None, 0, None, 0, None, "Abierto", "EJEMPLO – pendiente de datos reales"],
]
INPUT_COLS = [c for c, *_ in ECOLS if c not in ("A", "AB", "AC")]
for i in range(DATA0, DATAN + 1):
    put(en, f"A{i}", f'=IF(B{i}="","",ROW()-{DATA0 - 1})', border=True, align=CENTER)
    ej = EJ_EN[i - DATA0] if (i - DATA0 < len(EJ_EN) and not SIN_EJEMPLOS) else [None] * len(INPUT_COLS)
    for c, v in zip(INPUT_COLS, ej):
        f = next(e[3] for e in ECOLS if e[0] == c)
        inp(en, f"{c}{i}", v, f)
for c, h, g, f, w, v, com in ECOLS:
    rng = f"{c}{DATA0}:{c}{DATAN}"
    if v == "date":
        dv_date(en, rng)
    elif v == "num":
        dv_num(en, rng)
    elif v:
        dv_list(en, rng, v)
en.freeze_panes = "C6"
en.auto_filter.ref = f"A5:AC{DATAN}"

# ================================================================ CALCULO
ca = sheet("CALCULO", "548235")
title(ca, "Cálculo de la rentabilidad por contenedor (automático – no escribir)",
      "Cada fila corresponde a la misma fila de ENTRADA. Importes en USD. Los costos fijos se reparten entre los contenedores del mes.")

K_ING = "(1-Pct_Reclamaciones-Pct_Banco-Imp_Ventas_Serv-Contrib_Territorial)"
G = 'IF([id]="","",{})'.format
CSPEC = [
    # clave, encabezado, grupo, formato, fórmula
    ("n", "Nº", "IDENTIFICACIÓN", "int", G("ROW()-5")),
    ("id", "ID contenedor", "IDENTIFICACIÓN", "text", 'IF(E[B]="","",E[B])'),
    ("fecha", "Fecha de arribo", "IDENTIFICACIÓN", "date", G("E[C]")),
    ("mes", "Mes", "IDENTIFICACIÓN", "mon", G("DATE(YEAR([fecha]),MONTH([fecha]),1)")),
    ("trans", "Transitaria", "IDENTIFICACIÓN", "text", G('E[D]&""')),
    ("estado", "Estado", "IDENTIFICACIÓN", "text", G('E[Z]&""')),
    ("kgO", "kg Occidente", "VOLUMEN", "kg", G("N(E[G])")),
    ("kgC", "kg Centro", "VOLUMEN", "kg", G("N(E[H])")),
    ("kgR", "kg Oriente", "VOLUMEN", "kg", G("N(E[I])")),
    ("kgT", "kg total", "VOLUMEN", "kg", G("[kgO]+[kgC]+[kgR]")),
    ("vig", "Fila de precios vigente", "PRECIOS", "int",
     G(f"IFERROR(MATCH(SUMPRODUCT(MAX((PRECIOS!$A${PR0}:$A${PRN}<=[fecha])*PRECIOS!$A${PR0}:$A${PRN})),"
       f"PRECIOS!$A${PR0}:$A${PRN},0),1)")),
    ("vigD", "Precios vigentes desde", "PRECIOS", "date", G(f"INDEX(PRECIOS!$A${PR0}:$A${PRN},[vig])")),
    ("ingO", "Ingreso Occidente", "INGRESOS", "usd", G("[kgO]*P[B]")),
    ("ingC", "Ingreso Centro", "INGRESOS", "usd", G("[kgC]*P[C]")),
    ("ingR", "Ingreso Oriente", "INGRESOS", "usd", G("[kgR]*P[D]")),
    ("otrI", "Otros ingresos", "INGRESOS", "usd", G("N(E[X])")),
    ("ingT", "INGRESO TOTAL", "INGRESOS", "usd", G("[ingO]+[ingC]+[ingR]+[otrI]")),
    ("viajes", "Viajes al puerto", "ETAPA 1 · PUERTO", "num", G('IF(E[J]="",Std_Viajes_Puerto,E[J])')),
    ("cita", "Cita del puerto", "ETAPA 1 · PUERTO", "usd", G("[viajes]*P[F]")),
    ("litPs", "Litros al puerto (norma)", "ETAPA 1 · PUERTO", "num", G("[viajes]*Litros_Puerto_Viaje")),
    ("litP", "Litros al puerto (aplicados)", "ETAPA 1 · PUERTO", "num", G('IF(E[K]="",[litPs],E[K])')),
    ("combP", "Combustible al puerto", "ETAPA 1 · PUERTO", "usd", G("[litP]*P[E]")),
    ("dieta", "Dieta del chofer", "ETAPA 1 · PUERTO", "usd", G('IF(E[L]="",[viajes]*P[G],E[L])')),
    ("desg", "Desgaste del vehículo propio", "ETAPA 1 · PUERTO", "usd", G("[viajes]*Km_Puerto_Viaje*P[N]")),
    ("totP", "TOTAL PUERTO", "ETAPA 1 · PUERTO", "usd", G("[cita]+[combP]+[dieta]+[desg]")),
    ("jorTs", "Jornales de desagrupe (norma)", "ETAPA 2 · TRANSITARIA", "num",
     G("IF([kgT]=0,0,ROUNDUP([kgT]/Rend_Desagrupe,0))")),
    ("jorT", "Jornales de desagrupe (aplicados)", "ETAPA 2 · TRANSITARIA", "num", G('IF(E[M]="",[jorTs],E[M])')),
    ("moT", "Mano de obra del desagrupe", "ETAPA 2 · TRANSITARIA", "usd", G("[jorT]*P[H]")),
    ("servT", "Servicio de la transitaria", "ETAPA 2 · TRANSITARIA", "usd",
     G('IF([trans]="ENCI",P[J],IF([trans]="A.V",P[K],0))')),
    ("estad", "Estadía / almacenaje", "ETAPA 2 · TRANSITARIA", "usd", G("N(E[N])")),
    ("totT", "TOTAL TRANSITARIA", "ETAPA 2 · TRANSITARIA", "usd", G("[moT]+[servT]+[estad]")),
    ("viajO", "Viajes de distribución en Occidente", "OCCIDENTE", "num", G("IF([kgO]=0,0,ROUNDUP([kgO]/Cap_Camion_OCC,0))")),
    ("litDOs", "Litros de distribución (norma)", "OCCIDENTE", "num", G("[viajO]*L_Dist_OCC")),
    ("litDO", "Litros de distribución (aplicados)", "OCCIDENTE", "num", G('IF(E[O]="",[litDOs],E[O])')),
    ("combDO", "TOTAL OCCIDENTE (combustible)", "OCCIDENTE", "usd", G("[litDO]*P[E]")),
    ("viajC", "Viajes a Centro", "CENTRO", "num", G("IF([kgC]=0,0,ROUNDUP([kgC]/Cap_Camion_CEN,0))")),
    ("litTCs", "Litros de traslado (norma)", "CENTRO", "num", G("[viajC]*L_Tras_CEN")),
    ("litTC", "Litros de traslado (aplicados)", "CENTRO", "num", G('IF(E[P]="",[litTCs],E[P])')),
    ("litDCs", "Litros de distribución (norma)", "CENTRO", "num", G("[viajC]*L_Dist_CEN")),
    ("litDC", "Litros de distribución (aplicados)", "CENTRO", "num", G('IF(E[Q]="",[litDCs],E[Q])')),
    ("combC", "Combustible de Centro", "CENTRO", "usd", G("([litTC]+[litDC])*P[E]")),
    ("jorCs", "Jornales del centro transitorio (norma)", "CENTRO", "num", G("IF([kgC]=0,0,ROUNDUP([kgC]/Rend_Desagrupe_CT,0))")),
    ("jorC", "Jornales del centro transitorio (aplicados)", "CENTRO", "num", G('IF(E[R]="",[jorCs],E[R])')),
    ("moC", "Desagrupe en el centro transitorio", "CENTRO", "usd", G("[jorC]*P[I]")),
    ("totC", "TOTAL CENTRO", "CENTRO", "usd", G("[combC]+[moC]")),
    ("viajR", "Viajes a Oriente", "ORIENTE", "num", G("IF([kgR]=0,0,ROUNDUP([kgR]/Cap_Camion_ORI,0))")),
    ("litTRs", "Litros de traslado (norma)", "ORIENTE", "num", G("[viajR]*L_Tras_ORI")),
    ("litTR", "Litros de traslado (aplicados)", "ORIENTE", "num", G('IF(E[S]="",[litTRs],E[S])')),
    ("litDRs", "Litros de distribución (norma)", "ORIENTE", "num", G("[viajR]*L_Dist_ORI")),
    ("litDR", "Litros de distribución (aplicados)", "ORIENTE", "num", G('IF(E[T]="",[litDRs],E[T])')),
    ("combR", "Combustible de Oriente", "ORIENTE", "usd", G("([litTR]+[litDR])*P[E]")),
    ("jorRs", "Jornales del centro transitorio (norma)", "ORIENTE", "num", G("IF([kgR]=0,0,ROUNDUP([kgR]/Rend_Desagrupe_CT,0))")),
    ("jorR", "Jornales del centro transitorio (aplicados)", "ORIENTE", "num", G('IF(E[U]="",[jorRs],E[U])')),
    ("moR", "Desagrupe en el centro transitorio", "ORIENTE", "usd", G("[jorR]*P[I]")),
    ("totR", "TOTAL ORIENTE", "ORIENTE", "usd", G("[combR]+[moR]")),
    ("totD", "TOTAL DISTRIBUCIÓN", "ORIENTE", "usd", G("[combDO]+[totC]+[totR]")),
    ("comO", "Comisión Occidente", "COMERCIALIZACIÓN VARIABLE", "usd",
     G('IF(Base_Comision="% DEL INGRESO",[ingO]*P[L],[kgO]*P[M])')),
    ("comC", "Comisión Centro", "COMERCIALIZACIÓN VARIABLE", "usd",
     G('IF(Base_Comision="% DEL INGRESO",[ingC]*P[L],[kgC]*P[M])')),
    ("comR", "Comisión Oriente", "COMERCIALIZACIÓN VARIABLE", "usd",
     G('IF(Base_Comision="% DEL INGRESO",[ingR]*P[L],[kgR]*P[M])')),
    ("comT", "TOTAL COMISIÓN VARIABLE", "COMERCIALIZACIÓN VARIABLE", "usd", G("[comO]+[comC]+[comR]")),
    ("recl", "Provisión para reclamaciones", "OTROS VARIABLES", "usd", G("[ingT]*Pct_Reclamaciones")),
    ("banco", "Comisiones bancarias", "OTROS VARIABLES", "usd", G("[ingT]*Pct_Banco")),
    ("otrG", "Otros gastos directos", "OTROS VARIABLES", "usd", G("N(E[V])")),
    ("totV", "TOTAL COSTOS VARIABLES", "MARGEN DE CONTRIBUCIÓN", "usd",
     G("[totP]+[totT]+[totD]+[comT]+[recl]+[banco]+[otrG]")),
    ("MC", "MARGEN DE CONTRIBUCIÓN", "MARGEN DE CONTRIBUCIÓN", "usd", G("[ingT]-[totV]")),
    ("MCp", "MC %", "MARGEN DE CONTRIBUCIÓN", "pct", G("IF([ingT]=0,0,[MC]/[ingT])")),
    ("MCkg", "MC por kg", "MARGEN DE CONTRIBUCIÓN", "usdkg", G("IF([kgT]=0,0,[MC]/[kgT])")),
    ("colMes", "Columna del mes en COSTOS_FIJOS", "COSTOS FIJOS ASIGNADOS", "int",
     G(f"IFERROR(MATCH([mes],{CF_MONTHS},0),0)")),
    ("nMes", "Contenedores del mes", "COSTOS FIJOS ASIGNADOS", "int", G("COUNTIFS(COL[mes],[mes])")),
    ("kgMes", "kg del mes", "COSTOS FIJOS ASIGNADOS", "kg", G("SUMIFS(COL[kgT],COL[mes],[mes])")),
    ("shG", "% de los fijos generales", "COSTOS FIJOS ASIGNADOS", "pct",
     G('IF(Metodo_Prorrateo="KG",IF([kgMes]=0,0,[kgT]/[kgMes]),1/[nMes])')),
    ("fG", "Fijos generales asignados", "COSTOS FIJOS ASIGNADOS", "usd",
     G(f"IF([colMes]=0,0,INDEX({cf_row('GENERAL')},[colMes]))*[shG]")),
] + [
    item
    for k, pool in [("O", "OCC"), ("C", "CEN"), ("R", "ORI")]
    for item in [
        (f"sh{k}", f"% de los fijos {pool}", "COSTOS FIJOS ASIGNADOS", "pct",
         G(f'IF(Metodo_Prorrateo="KG",IF(SUMIFS(COL[kg{k}],COL[mes],[mes])=0,0,[kg{k}]/SUMIFS(COL[kg{k}],COL[mes],[mes])),'
           f'IF([kg{k}]>0,1/COUNTIFS(COL[mes],[mes],COL[kg{k}],">0"),0))')),
        (f"f{k}", f"Fijos {pool} asignados", "COSTOS FIJOS ASIGNADOS", "usd",
         G(f"IF([colMes]=0,0,INDEX({cf_row(pool)},[colMes]))*[sh{k}]")),
    ]
] + [
    ("totF", "TOTAL FIJOS ASIGNADOS", "COSTOS FIJOS ASIGNADOS", "usd", G("[fG]+[fO]+[fC]+[fR]")),
    ("ivs", "Impuesto sobre ventas y servicios", "TRIBUTOS", "usd", G("[ingT]*Imp_Ventas_Serv")),
    ("terr", "Contribución territorial", "TRIBUTOS", "usd", G("[ingT]*Contrib_Territorial")),
    ("UAI", "UTILIDAD ANTES DE IMPUESTO SOBRE UTILIDADES", "RESULTADO", "usd", G("[MC]-[totF]-[ivs]-[terr]")),
    ("iu", "Impuesto sobre utilidades (estimado)", "RESULTADO", "usd", G("MAX(0,[UAI])*Imp_Utilidades")),
    ("UN", "UTILIDAD NETA ESTIMADA", "RESULTADO", "usd", G("[UAI]-[iu]")),
    ("MN", "Margen neto %", "RESULTADO", "pct", G("IF([ingT]=0,0,[UN]/[ingT])")),
    ("ctKg", "Costo total por kg", "RESULTADO", "usdkg", G("IF([kgT]=0,0,([totV]+[totF]+[ivs]+[terr])/[kgT])")),
    ("ingKg", "Ingreso por kg", "RESULTADO", "usdkg", G("IF([kgT]=0,0,[ingT]/[kgT])")),
    ("uaiKg", "Utilidad antes de impuesto por kg", "RESULTADO", "usdkg", G("IF([kgT]=0,0,[UAI]/[kgT])")),
    ("shared", "Costos compartidos (puerto + transitaria + otros directos + fijos generales)", "RESULTADO POR REGIÓN", "usd",
     G("[totP]+[totT]+[otrG]+[fG]")),
    ("uO", "Utilidad Occidente", "RESULTADO POR REGIÓN", "usd",
     G(f"IF([kgT]=0,0,([ingO]+[otrI]*[kgO]/[kgT])*{K_ING}-[combDO]-[comO]-[fO]-[shared]*[kgO]/[kgT])")),
    ("uC", "Utilidad Centro", "RESULTADO POR REGIÓN", "usd",
     G(f"IF([kgT]=0,0,([ingC]+[otrI]*[kgC]/[kgT])*{K_ING}-[totC]-[comC]-[fC]-[shared]*[kgC]/[kgT])")),
    ("uR", "Utilidad Oriente", "RESULTADO POR REGIÓN", "usd",
     G(f"IF([kgT]=0,0,([ingR]+[otrI]*[kgR]/[kgT])*{K_ING}-[totR]-[comR]-[fR]-[shared]*[kgR]/[kgT])")),
    ("chk", "Control: diferencia de cuadre (debe ser 0)", "RESULTADO POR REGIÓN", "usd",
     G("IF([kgT]=0,0,ROUND([UAI]-[uO]-[uC]-[uR],4))")),
    ("litS", "Litros totales según norma", "CONTROL", "num", G("[litPs]+[litDOs]+[litTCs]+[litDCs]+[litTRs]+[litDRs]")),
    ("litR", "Litros totales aplicados", "CONTROL", "num", G("[litP]+[litDO]+[litTC]+[litDC]+[litTR]+[litDR]")),
    ("desv", "Desvío del combustible frente a la norma", "CONTROL", "pct", G("IF([litS]=0,0,[litR]/[litS]-1)")),
    ("ciclo", "Días de ciclo (arribo → fin)", "CONTROL", "int", G('IF(E[Y]="","",E[Y]-[fecha])')),
    ("alert", "ALERTAS", "CONTROL", "text",
     G('IF(E[C]="","Falta fecha. ","")&IF([colMes]=0,"Mes fuera del rango de COSTOS_FIJOS. ","")'
       '&IF([kgT]=0,"Sin kg. ","")&IF([trans]="","Falta transitaria. ","")'
       '&IF(AND([litPs]>0,[litP]/[litPs]-1>Tolerancia_Comb),"Combustible al puerto "&TEXT([litP]/[litPs]-1,"+0%")&" sobre norma. ","")'
       '&IF(AND([litDOs]>0,[litDO]/[litDOs]-1>Tolerancia_Comb),"Combustible distribución Occidente "&TEXT([litDO]/[litDOs]-1,"+0%")&" sobre norma. ","")'
       '&IF(AND([litTCs]>0,[litTC]/[litTCs]-1>Tolerancia_Comb),"Combustible traslado Centro "&TEXT([litTC]/[litTCs]-1,"+0%")&" sobre norma. ","")'
       '&IF(AND([litDCs]>0,[litDC]/[litDCs]-1>Tolerancia_Comb),"Combustible distribución Centro "&TEXT([litDC]/[litDCs]-1,"+0%")&" sobre norma. ","")'
       '&IF(AND([litTRs]>0,[litTR]/[litTRs]-1>Tolerancia_Comb),"Combustible traslado Oriente "&TEXT([litTR]/[litTRs]-1,"+0%")&" sobre norma. ","")'
       '&IF(AND([litDRs]>0,[litDR]/[litDRs]-1>Tolerancia_Comb),"Combustible distribución Oriente "&TEXT([litDR]/[litDRs]-1,"+0%")&" sobre norma. ","")'
       f'&IF(AND(E[C]<>"",[fecha]<PRECIOS!$A${PR0}),"Fecha anterior a la 1ª vigencia de precios. ","")'
       '&IF([UAI]<0,"Contenedor con pérdida. ","")')),
]
CK = {}  # clave -> letra
for i, (k, *_r) in enumerate(CSPEC):
    CK[k] = CL(i + 1)
assert len(CK) == len(CSPEC), "claves duplicadas"


def render(tpl, r):
    s = re.sub(r"E\[([A-Z]+)\]", lambda m: f"ENTRADA!${m.group(1)}{r}", tpl)
    s = re.sub(r"P\[([A-Z]+)\]", lambda m: f"INDEX(PRECIOS!${m.group(1)}${PR0}:${m.group(1)}${PRN},{CK['vig']}{r})", s)
    s = re.sub(r"COL\[(\w+)\]", lambda m: f"${CK[m.group(1)]}${DATA0}:${CK[m.group(1)]}${DATAN}", s)
    s = re.sub(r"\[(\w+)\]", lambda m: f"{CK[m.group(1)]}{r}", s)
    return "=" + s


TOTAL_KEYS = {"ingT", "totP", "totT", "totD", "comT", "totV", "MC", "totF", "UAI", "UN", "kgT"}
grp_cols = {}
for k, h, g, f, tpl in CSPEC:
    grp_cols.setdefault(g, []).append(CK[k])
for g, cols in grp_cols.items():
    a, b = cols[0], cols[-1]
    if a != b:
        ca.merge_cells(f"{a}4:{b}4")
    hdr(ca, f"{a}4", g, FILL_GRP, F_BOLD)
for k, h, g, f, tpl in CSPEC:
    L = CK[k]
    hdr(ca, f"{L}5", h)
    ca.column_dimensions[L].width = 34 if k == "alert" else (16 if k == "id" else 12)
    bold = k in TOTAL_KEYS
    link = tpl.startswith('IF(E[B]') or tpl in (G("E[C]"),)
    for r in range(DATA0, DATAN + 1):
        c = ca[f"{L}{r}"]
        c.value = render(tpl, r)
        c.font = F_LINK if link else (F_BOLD if bold else F_BASE)
        c.number_format = FMT[f]
        if bold:
            c.fill = FILL_TOTAL
ca.row_dimensions[5].height = 66
ca.freeze_panes = "C6"
ca.auto_filter.ref = f"A5:{CK['alert']}{DATAN}"
ca.conditional_formatting.add(
    f"{CK['alert']}{DATA0}:{CK['alert']}{DATAN}",
    FormulaRule(formula=[f'LEN({CK["alert"]}{DATA0})>0'], fill=PatternFill("solid", fgColor="FCE4D6")))
protect(ca)


def CR(key):
    L = CK[key]
    return f"CALCULO!${L}${DATA0}:${L}${DATAN}"


# resultados visibles en ENTRADA
for r in range(DATA0, DATAN + 1):
    put(en, f"AB{r}", f"=CALCULO!{CK['alert']}{r}", F_LINK, border=True)
    put(en, f"AC{r}", f"=CALCULO!{CK['UAI']}{r}", F_LINK, "usd", border=True)
en.conditional_formatting.add(f"AB{DATA0}:AB{DATAN}",
                              FormulaRule(formula=[f"LEN(AB{DATA0})>0"], fill=PatternFill("solid", fgColor="FCE4D6")))
protect(en)
en.protection.insertRows = True
en.protection.deleteRows = True

# ================================================================ COMERCIAL (contenido)
title(co, "Comercialización: 21 gestores y liquidación mensual",
      "El pago fijo de todos los gestores pasa a COSTOS_FIJOS. La parte variable se liquida cada mes y se concilia "
      "con los kg o ingresos reales del mes.")
put(co, "A4", "Mes a liquidar", F_BOLD)
inp(co, "C4", dt.date(2026, 10, 1), "mon")
dv_date(co, "C4")
put(co, "A5", "Vigencia de precios del mes", F_BOLD)
put(co, "C5", f"=IFERROR(MATCH(SUMPRODUCT(MAX((PRECIOS!$A${PR0}:$A${PRN}<=EOMONTH(C4,0))*PRECIOS!$A${PR0}:$A${PRN})),"
              f"PRECIOS!$A${PR0}:$A${PRN},0),1)", fmt="int")
put(co, "A6", "Base de la comisión", F_BOLD)
put(co, "C6", "=Base_Comision", F_LINK)
put(co, "A7", "Tasa variable aplicable", F_BOLD)
put(co, "C7", f'=IF(C6="% DEL INGRESO",INDEX(PRECIOS!$L${PR0}:$L${PRN},C5),INDEX(PRECIOS!$M${PR0}:$M${PRN},C5))',
    fmt='0.0000')
put(co, "D7", '=IF(C6="% DEL INGRESO","de lo cobrado (USD)","USD por kg entregado")', F_SUB)
put(co, "A8", "Pago fijo mensual por defecto (USD)", F_BOLD)
inp(co, "C8", 50, "usd")
put(co, "D8", "ILUSTRATIVO. Todos los gestores cobran lo mismo; se puede ajustar por persona en la columna F.", F_SUB)
COCOLS = [("A", "Nº", 5), ("B", "Nombre del gestor", 26), ("C", "Grupo", 15), ("D", "Territorio", 20),
          ("E", "Región", 11), ("F", "Pago fijo mensual (USD)", 13),
          ("G", "Base variable del mes (USD cobrados o kg)", 16), ("H", "Pago variable (USD)", 13),
          ("I", "TOTAL A PAGAR (USD)", 14)]
for c, t, w in COCOLS:
    hdr(co, f"{c}{CO0 - 1}", t)
    co.column_dimensions[c].width = w
co.row_dimensions[CO0 - 1].height = 45
roster = [("Grupo central", "Oficina central")] * 5 + [("Provincial", t) for t, _ in TERRITORIOS[:16]]
# base variable de ejemplo para oct-2026 (ingresos de los contenedores 3 y 4 por región, repartidos)
peso = {"Pinar del Río": 1, "Artemisa": 1, "La Habana": 4, "Mayabeque": 1, "Matanzas": 2, "Isla de la Juventud": 0.5,
        "Cienfuegos": 1, "Villa Clara": 2, "Sancti Spíritus": 1, "Ciego de Ávila": 1, "Camagüey": 2,
        "Las Tunas": 1, "Holguín": 2, "Granma": 1.5, "Santiago de Cuba": 2.5, "Guantánamo": 1}
ing_reg = {"Occidente": (6100 + 3800) * 0.8, "Centro": (2900 + 1500) * 0.8, "Oriente": (2400 + 3900) * 0.8}
base_ej = {}
for reg, total in ing_reg.items():
    ts = [t for t, rg in TERRITORIOS[:16] if rg == reg]
    tot_w = sum(peso[t] for t in ts)
    acum = 0
    for t in ts[:-1]:
        v = round(total * peso[t] / tot_w, 2)
        base_ej[t] = v
        acum += v
    base_ej[ts[-1]] = round(total - acum, 2)
for i, (grp, terr) in enumerate(roster):
    rr = CO0 + i
    put(co, f"A{rr}", i + 1, border=True, align=CENTER)
    inp(co, f"B{rr}", f"Gestor {i + 1:02d} (editar)")
    inp(co, f"C{rr}", grp)
    inp(co, f"D{rr}", terr)
    put(co, f"E{rr}", f'=IFERROR(INDEX(L_TerrRegion,MATCH(D{rr},L_Territorio,0)),"")', border=True)
    inp(co, f"F{rr}", "=$C$8", "usd")
    inp(co, f"G{rr}", base_ej.get(terr) if grp == "Provincial" else 0, "usd")
    put(co, f"H{rr}", f"=N(G{rr})*$C$7", fmt="usd", border=True)
    put(co, f"I{rr}", f"=N(F{rr})+H{rr}", F_BOLD, "usd", border=True)
put(co, f"B{CO_TOT}", "TOTAL", F_BOLD, fill=FILL_TOTAL, border=True)
for c in "FGHI":
    put(co, f"{c}{CO_TOT}", f"=SUM({c}{CO0}:{c}{CON})", F_BOLD, "usd", FILL_TOTAL, border=True)
dv_list(co, f"C{CO0}:C{CON}", "L_Grupo")
dv_list(co, f"D{CO0}:D{CON}", "L_Territorio")
dv_num(co, f"F{CO0}:G{CON}")
rq = CO_TOT + 2
put(co, f"A{rq}", "CONCILIACIÓN DEL MES", F_H2)
put(co, f"A{rq + 1}", "Base variable declarada por los gestores", border=True)
put(co, f"G{rq + 1}", f"=G{CO_TOT}", fmt="usd", border=True)
put(co, f"A{rq + 2}", "Base real del mes según las operaciones (CALCULO)", border=True)
put(co, f"G{rq + 2}",
    f'=IF(C6="% DEL INGRESO",SUMIFS({CR("ingO")},{CR("mes")},C4)+SUMIFS({CR("ingC")},{CR("mes")},C4)'
    f'+SUMIFS({CR("ingR")},{CR("mes")},C4),SUMIFS({CR("kgT")},{CR("mes")},C4))', F_LINK, "usd", border=True)
put(co, f"A{rq + 3}", "Diferencia", F_BOLD, border=True)
put(co, f"G{rq + 3}", f"=G{rq + 1}-G{rq + 2}", F_BOLD, "usd", border=True)
put(co, f"A{rq + 4}", "Estado de la conciliación", F_BOLD, border=True)
put(co, f"G{rq + 4}", f'=IF(ABS(G{rq + 3})<0.01,"CUADRA","REVISAR: la base declarada no coincide con las operaciones")', F_BOLD)
put(co, f"A{rq + 5}", "Pago variable calculado en el modelo de costos (CALCULO)", border=True)
put(co, f"G{rq + 5}", f"=SUMIFS({CR('comT')},{CR('mes')},C4)", F_LINK, "usd", border=True)
put(co, f"A{rq + 7}", "Nota: los ingresos de la conciliación se calculan a la tarifa por kg. Si se liquida sobre lo efectivamente COBRADO "
                      "(recomendado), la diferencia muestra lo pendiente de cobro.", F_SUB)
co.freeze_panes = "A12"
protect(co)
CO_CONC = rq + 4

# ================================================================ RESUMEN_MES
rm = sheet("RESUMEN_MES", "548235")
title(rm, "Resumen mensual de rentabilidad (automático)",
      "Utilidad del mes = margen de contribución − TODOS los costos fijos del mes − tributos sobre ingresos.")
RCOLS = [
    ("A", "Mes", "mon"), ("B", "Contenedores", "int"), ("C", "kg Occidente", "kg"), ("D", "kg Centro", "kg"),
    ("E", "kg Oriente", "kg"), ("F", "kg total", "kg"), ("G", "Ingresos", "usd0"), ("H", "Costos variables", "usd0"),
    ("I", "Margen de contribución", "usd0"), ("J", "MC %", "pct"), ("K", "Costos fijos del mes (hasta el corte)", "usd0"),
    ("L", "Fijos absorbidos por los contenedores", "usd0"), ("M", "Fijos no absorbidos", "usd0"),
    ("N", "Tributos sobre ingresos", "usd0"), ("O", "UTILIDAD ANTES DE IMPUESTO", "usd0"), ("P", "Margen %", "pct"),
    ("Q", "Costo total por kg", "usdkg"), ("R", "Ingreso por kg", "usdkg"), ("S", "Utilidad por kg", "usdkg"),
    ("T", "Punto de equilibrio (kg)", "kg"), ("U", "Punto de equilibrio (contenedores)", "num"),
    ("V", "Margen de seguridad", "pct"), ("W", "Utilidad Occidente", "usd0"), ("X", "Utilidad Centro", "usd0"),
    ("Y", "Utilidad Oriente", "usd0"), ("Z", "Ajuste por fijos no absorbidos", "usd0"),
    ("AA", "Control (debe ser 0)", "usd"), ("AB", "Costos totales", "usd0"), ("AC", "Desvío medio del combustible", "pct"),
]
for c, t, f in RCOLS:
    hdr(rm, f"{c}5", t)
    rm.column_dimensions[c].width = 12
rm.column_dimensions["A"].width = 10
rm.row_dimensions[5].height = 54
R0, RN = 6, 6 + NMES - 1
put(rm, "A3", "Mes de corte aplicado:", F_BOLD)
put(rm, "C3", f'=IF(Mes_Corte="",IF(COUNT({CR("mes")})=0,Fecha_Inicio,MAX({CR("mes")})),DATE(YEAR(Mes_Corte),MONTH(Mes_Corte),1))',
    F_BOLD, "mon")
put(rm, "D3", "Los meses posteriores al corte no cargan costos fijos (aún no han ocurrido). Ajustable en PARAMETROS > Mes_Corte.", F_SUB)


def S(key, r):
    return f"SUMIFS({CR(key)},{CR('mes')},$A{r})"


for r in range(R0, RN + 1):
    k = r - R0
    F = {
        "A": f"=EDATE(Fecha_Inicio,{k})",
        "B": f"=COUNTIFS({CR('mes')},$A{r})",
        "C": "=" + S("kgO", r), "D": "=" + S("kgC", r), "E": "=" + S("kgR", r),
        "F": f"=C{r}+D{r}+E{r}",
        "G": "=" + S("ingT", r), "H": "=" + S("totV", r),
        "I": f"=G{r}-H{r}", "J": f"=IF(G{r}=0,0,I{r}/G{r})",
        "K": f"=IF($A{r}>$C$3,0,INDEX({cf_row('ALL')},{k + 1}))",
        "L": "=" + S("totF", r), "M": f"=K{r}-L{r}",
        "N": f"={S('ivs', r)}+{S('terr', r)}",
        "O": f"=I{r}-K{r}-N{r}", "P": f"=IF(G{r}=0,0,O{r}/G{r})",
        "Q": f"=IF(F{r}=0,0,(H{r}+K{r}+N{r})/F{r})", "R": f"=IF(F{r}=0,0,G{r}/F{r})",
        "S": f"=IF(F{r}=0,0,O{r}/F{r})",
        "T": f'=IF(F{r}=0,"s/d",IF(I{r}-N{r}<=0,"No cubre",K{r}/((I{r}-N{r})/F{r})))',
        "U": f'=IF(ISNUMBER(T{r}),T{r}/(F{r}/B{r}),"s/d")',
        "V": f'=IF(ISNUMBER(T{r}),(F{r}-T{r})/F{r},"s/d")',
        "W": "=" + S("uO", r), "X": "=" + S("uC", r), "Y": "=" + S("uR", r),
        "Z": f"=-M{r}", "AA": f"=ROUND(O{r}-(W{r}+X{r}+Y{r}+Z{r}),4)",
        "AB": f"=H{r}+K{r}+N{r}",
        "AC": f"=IFERROR(SUMIFS({CR('litR')},{CR('mes')},$A{r})/SUMIFS({CR('litS')},{CR('mes')},$A{r})-1,0)",
    }
    for c, t, f in RCOLS:
        put(rm, f"{c}{r}", F[c], F_BOLD if c in ("O",) else F_BASE, f, border=True)
RT = RN + 1
put(rm, f"A{RT}", "TOTAL", F_BOLD, fill=FILL_TOTAL, border=True)
for c, t, f in RCOLS[1:]:
    if c in ("J",):
        v = f"=IF(G{RT}=0,0,I{RT}/G{RT})"
    elif c == "P":
        v = f"=IF(G{RT}=0,0,O{RT}/G{RT})"
    elif c == "Q":
        v = f"=IF(F{RT}=0,0,(H{RT}+K{RT}+N{RT})/F{RT})"
    elif c == "R":
        v = f"=IF(F{RT}=0,0,G{RT}/F{RT})"
    elif c == "S":
        v = f"=IF(F{RT}=0,0,O{RT}/F{RT})"
    elif c in ("T", "U", "V", "AC"):
        v = None
    else:
        v = f"=SUM({c}{R0}:{c}{RN})"
    put(rm, f"{c}{RT}", v, F_BOLD, f, FILL_TOTAL, border=True)
# tabla anual
YA = RT + 3
put(rm, f"A{YA - 1}", "LIQUIDACIÓN ANUAL ESTIMADA DEL IMPUESTO SOBRE UTILIDADES", F_H2)
for c, t in zip("ABCDE", ["Año", "Utilidad antes de impuesto", "Impuesto sobre utilidades (estimado)", "Utilidad neta estimada", "Margen neto %"]):
    hdr(rm, f"{c}{YA}", t)
rm.row_dimensions[YA].height = 42
for j in range(3):
    rr = YA + 1 + j
    put(rm, f"A{rr}", f"=YEAR(Fecha_Inicio)+{j}", fmt="0", border=True)
    put(rm, f"B{rr}", f"=SUMPRODUCT((YEAR($A${R0}:$A${RN})=A{rr})*$O${R0}:$O${RN})", fmt="usd0", border=True)
    put(rm, f"C{rr}", f"=MAX(0,B{rr})*Imp_Utilidades", fmt="usd0", border=True)
    put(rm, f"D{rr}", f"=B{rr}-C{rr}", F_BOLD, "usd0", border=True)
    put(rm, f"E{rr}", f"=IFERROR(D{rr}/SUMPRODUCT((YEAR($A${R0}:$A${RN})=A{rr})*$G${R0}:$G${RN}),0)", fmt="pct", border=True)
put(rm, f"A{YA + 5}", "El impuesto sobre utilidades se liquida por año fiscal, con pagos a cuenta. Las pérdidas de un mes compensan "
                      "las ganancias de otro dentro del mismo año. Confirme con el contador.", F_SUB)
rm.freeze_panes = "B6"
protect(rm)


def RS(c):
    return f"RESUMEN_MES!${c}${R0}:${c}${RN}"


# ================================================================ DASHBOARD
db = sheet("DASHBOARD", "1F3864")
title(db, "Panel de rentabilidad", None)
put(db, "A2", '="Empresa: "&Empresa', F_SUB)
widths(db, {"A": 2, "B": 30, "C": 15, "D": 13, "E": 13, "F": 3, "G": 30, "H": 15, "I": 13, "J": 13, "K": 13, "L": 13})
put(db, "B4", "Período desde (mes)", F_BOLD)
inp(db, "C4", dt.date(2026, 9, 1), "mon")
put(db, "B5", "Período hasta (mes)", F_BOLD)
inp(db, "C5", dt.date(2026, 10, 1), "mon")
put(db, "D4", "Elija el primer día del mes (dd/mm/aaaa) o seleccione de la lista.", F_SUB)
put(db, "D5", '=IF(C5<C4,"ATENCIÓN: «hasta» es anterior a «desde»","")', Font(name=FONT, size=10, bold=True, color="C00000"))
for ref in ("C4", "C5"):
    dv = DataValidation(type="list", formula1=f"={RS('A')}", allow_blank=False)
    db.add_data_validation(dv)
    dv.add(ref)
PER = f'{RS("A")},">="&$C$4,{RS("A")},"<="&$C$5'
PERC = f'{CR("mes")},">="&$C$4,{CR("mes")},"<="&$C$5'


def RSUM(c):
    return f"SUMIFS({RS(c)},{PER})"


def CSUM(k):
    return f"SUMIFS({CR(k)},{PERC})"


hdr(db, "B7", "INDICADOR")
hdr(db, "C7", "VALOR")
hdr(db, "G7", "INDICADOR")
hdr(db, "H7", "VALOR")
KPI_L = [
    ("Contenedores", f"={RSUM('B')}", "int"),
    ("kg distribuidos", f"={RSUM('F')}", "kg"),
    ("Ingresos", f"={RSUM('G')}", "usd0"),
    ("Costos variables", f"={RSUM('H')}", "usd0"),
    ("Margen de contribución", "=C10-C11", "usd0"),
    ("MC %", "=IF(C10=0,0,C12/C10)", "pct"),
    ("Costos fijos del período", f"={RSUM('K')}", "usd0"),
    ("Tributos sobre ingresos", f"={RSUM('N')}", "usd0"),
    ("UTILIDAD ANTES DE IMPUESTO", "=C12-C14-C15", "usd0"),
    ("Impuesto sobre utilidades (estimado)", "=MAX(0,C16)*Imp_Utilidades", "usd0"),
    ("UTILIDAD NETA ESTIMADA", "=C16-C17", "usd0"),
]
KPI_R = [
    ("Margen neto %", "=IF(C10=0,0,C18/C10)", "pct"),
    ("Ingreso por kg", "=IF(C9=0,0,C10/C9)", "usdkg"),
    ("Costo total por kg", "=IF(C9=0,0,(C11+C14+C15)/C9)", "usdkg"),
    ("Utilidad antes de impuesto por kg", "=IF(C9=0,0,C16/C9)", "usdkg"),
    ("Utilidad media por contenedor", "=IF(C8=0,0,C16/C8)", "usd0"),
    ("Punto de equilibrio del período (kg)", '=IF(C9=0,"s/d",IF(C12-C15<=0,"No cubre",C14/((C12-C15)/C9)))', "kg"),
    ("Punto de equilibrio (contenedores)", '=IF(ISNUMBER(H13),H13/(C9/C8),"s/d")', "num"),
    ("Margen de seguridad", '=IF(ISNUMBER(H13),(C9-H13)/C9,"s/d")', "pct"),
    ("Fijos no absorbidos (capacidad ociosa)", f"={RSUM('M')}", "usd0"),
    ("Contenedores con alerta (todo el libro)", f'=COUNTIF({CR("alert")},"?*")', "int"),
    ("Desvío medio del combustible", f"=IFERROR({CSUM('litR')}/{CSUM('litS')}-1,0)", "pct"),
]
for i, (lab, f, fm) in enumerate(KPI_L):
    rr = 8 + i
    strong = lab.isupper()
    put(db, f"B{rr}", lab, F_BOLD if strong else F_BASE, fill=FILL_KPI, border=True)
    put(db, f"C{rr}", f, Font(name=FONT, size=11, bold=True), fm, FILL_KPI, border=True)
for i, (lab, f, fm) in enumerate(KPI_R):
    rr = 8 + i
    put(db, f"G{rr}", lab, fill=FILL_KPI, border=True)
    put(db, f"H{rr}", f, Font(name=FONT, size=11, bold=True), fm, FILL_KPI, border=True)

# estructura de costos
SR = 21
put(db, f"B{SR - 1}", "Estructura de costos del período", F_H2)
for c, t in zip("BCDE", ["Concepto", "USD", "% de los ingresos", "USD por kg"]):
    hdr(db, f"{c}{SR}", t)
EST = [
    ("Puerto (vehículo propio)", CSUM("totP")),
    ("Transitaria (desagrupe y estadía)", CSUM("totT")),
    ("Distribución Occidente", CSUM("combDO")),
    ("Distribución Centro", CSUM("totC")),
    ("Distribución Oriente", CSUM("totR")),
    ("Comercialización variable", CSUM("comT")),
    ("Reclamaciones y banca", f"{CSUM('recl')}+{CSUM('banco')}"),
    ("Otros gastos directos", CSUM("otrG")),
    ("Costos fijos", RSUM("K")),
    ("Tributos sobre ingresos", RSUM("N")),
]
for i, (lab, f) in enumerate(EST):
    rr = SR + 1 + i
    put(db, f"B{rr}", lab, border=True)
    put(db, f"C{rr}", "=" + f, fmt="usd0", border=True)
    put(db, f"D{rr}", f"=IF($C$10=0,0,C{rr}/$C$10)", fmt="pct", border=True)
    put(db, f"E{rr}", f"=IF($C$9=0,0,C{rr}/$C$9)", fmt="usdkg", border=True)
ER = SR + 1 + len(EST)
put(db, f"B{ER}", "TOTAL COSTOS", F_BOLD, fill=FILL_TOTAL, border=True)
put(db, f"C{ER}", f"=SUM(C{SR + 1}:C{ER - 1})", F_BOLD, "usd0", FILL_TOTAL, border=True)
put(db, f"D{ER}", f"=IF($C$10=0,0,C{ER}/$C$10)", F_BOLD, "pct", FILL_TOTAL, border=True)
put(db, f"E{ER}", f"=IF($C$9=0,0,C{ER}/$C$9)", F_BOLD, "usdkg", FILL_TOTAL, border=True)
put(db, f"B{ER + 1}", "Control: ingresos − costos − utilidad (debe ser 0)", F_SUB)
put(db, f"C{ER + 1}", f"=ROUND(C10-C{ER}-C16,2)", fmt="usd")

# por región
for c, t in zip("GHIJK", ["Región", "kg", "Ingresos", "Utilidad antes de impuesto", "Margen %"]):
    hdr(db, f"{c}{SR}", t)
put(db, f"G{SR - 1}", "Rentabilidad por región del período", F_H2)
REG = [("Occidente", "kgO", "ingO", "W"), ("Centro", "kgC", "ingC", "X"), ("Oriente", "kgR", "ingR", "Y")]
for i, (lab, kk, ik, uc) in enumerate(REG):
    rr = SR + 1 + i
    put(db, f"G{rr}", lab, border=True)
    put(db, f"H{rr}", "=" + CSUM(kk), fmt="kg", border=True)
    put(db, f"I{rr}", "=" + CSUM(ik), fmt="usd0", border=True)
    put(db, f"J{rr}", "=" + RSUM(uc), fmt="usd0", border=True)
    put(db, f"K{rr}", f"=IF(I{rr}=0,0,J{rr}/I{rr})", fmt="pct", border=True)
rr = SR + 4
put(db, f"G{rr}", "Fijos no absorbidos", border=True)
put(db, f"J{rr}", "=" + RSUM("Z"), fmt="usd0", border=True)
put(db, f"G{rr + 1}", "TOTAL", F_BOLD, fill=FILL_TOTAL, border=True)
put(db, f"H{rr + 1}", f"=SUM(H{SR + 1}:H{SR + 3})", F_BOLD, "kg", FILL_TOTAL, border=True)
put(db, f"I{rr + 1}", f"=SUM(I{SR + 1}:I{SR + 3})", F_BOLD, "usd0", FILL_TOTAL, border=True)
put(db, f"J{rr + 1}", f"=SUM(J{SR + 1}:J{rr})", F_BOLD, "usd0", FILL_TOTAL, border=True)
put(db, f"K{rr + 1}", f"=IF(I{rr + 1}=0,0,J{rr + 1}/I{rr + 1})", F_BOLD, "pct", FILL_TOTAL, border=True)
put(db, f"G{rr + 2}", "Los ingresos regionales no incluyen «otros ingresos», que se reparten en la utilidad según los kg.", F_SUB)

# tabla mensual para gráficos (hasta 12 meses del período)
MT = ER + 4
put(db, f"B{MT - 1}", "Evolución mensual del período (máx. 12 meses)", F_H2)
for c, t in zip("BCDE", ["Mes", "Ingresos", "Costos totales", "Utilidad antes de impuesto"]):
    hdr(db, f"{c}{MT}", t)
for i in range(12):
    rr = MT + 1 + i
    put(db, f"B{rr}", f'=IF(EDATE($C$4,{i})>$C$5,"",EDATE($C$4,{i}))', fmt="mmm-yy", border=True)
    for c, src in zip("CDE", ["G", "AB", "O"]):
        put(db, f"{c}{rr}", f'=IF($B{rr}="","",SUMIFS({RS(src)},{RS("A")},$B{rr}))', fmt="usd0", border=True)

BLUE, ORANGE, AQUA = "2A78D6", "EB6834", "1BAF7A"


def style_chart(ch, title_txt, ytitle):
    ch.title = title_txt
    ch.y_axis.title = ytitle
    ch.y_axis.majorGridlines.spPr = None
    ch.height, ch.width = 7.5, 16
    ch.legend.position = "b"
    ch.y_axis.delete = False
    ch.x_axis.delete = False


bc = BarChart()
bc.type = "col"
bc.grouping = "clustered"
data = Reference(db, min_col=3, max_col=4, min_row=MT, max_row=MT + 12)
cats = Reference(db, min_col=2, min_row=MT + 1, max_row=MT + 12)
bc.add_data(data, titles_from_data=True)
bc.set_categories(cats)
for s, colr in zip(bc.series, [BLUE, ORANGE]):
    s.graphicalProperties.solidFill = colr
    s.graphicalProperties.line.noFill = True
lc = LineChart()
lc.add_data(Reference(db, min_col=5, min_row=MT, max_row=MT + 12), titles_from_data=True)
lc.series[0].graphicalProperties.line.solidFill = AQUA
lc.series[0].graphicalProperties.line.width = 28000
lc.series[0].marker.symbol = "circle"
lc.series[0].marker.size = 7
lc.series[0].marker.graphicalProperties.solidFill = AQUA
style_chart(bc, "Ingresos, costos y utilidad por mes (USD)", "USD")
bc.y_axis.number_format = "#,##0"
bc += lc
db.add_chart(bc, "G35")

cs = BarChart()
cs.type = "bar"
cs.add_data(Reference(db, min_col=3, min_row=SR, max_row=ER - 1), titles_from_data=True)
cs.set_categories(Reference(db, min_col=2, min_row=SR + 1, max_row=ER - 1))
cs.series[0].graphicalProperties.solidFill = BLUE
style_chart(cs, "Estructura de costos del período (USD)", "USD")
cs.legend = None
cs.y_axis.number_format = "#,##0"
cs.gapWidth = 40
cs.height, cs.width = 10, 16
db.add_chart(cs, "B51")

rc = BarChart()
rc.type = "col"
rc.add_data(Reference(db, min_col=10, min_row=SR, max_row=SR + 3), titles_from_data=True)
rc.set_categories(Reference(db, min_col=7, min_row=SR + 1, max_row=SR + 3))
rc.series[0].graphicalProperties.solidFill = BLUE
style_chart(rc, "Utilidad antes de impuesto por región (USD)", "USD")
rc.legend = None
rc.y_axis.number_format = "#,##0"
rc.width, rc.height = 13, 10
db.add_chart(rc, "H51")
protect(db)

# ================================================================ SIMULADOR
sm = sheet("SIMULADOR", "7030A0")
title(sm, "Simulador de escenarios («¿qué pasaría si...?»)",
      "«Valor actual» trae los datos reales del libro. Escriba en «Valor a simular» para probar un escenario; bórrelo para volver al dato real. "
      "El simulador no modifica el libro.")
widths(sm, {"A": 3, "B": 44, "C": 14, "D": 14, "E": 14, "F": 16, "G": 13, "H": 13, "I": 13, "J": 13, "K": 13})
for c, t in zip("BCDEF", ["Supuesto", "Valor actual", "Valor a simular", "Valor usado", "Unidad"]):
    hdr(sm, f"{c}4", t)
LAST = f"MATCH(MAX(PRECIOS!$A${PR0}:$A${PRN}),PRECIOS!$A${PR0}:$A${PRN},0)"


def lastp(c):
    return f"INDEX(PRECIOS!${c}${PR0}:${c}${PRN},{LAST})"


REFMES = f"IFERROR(MATCH(MAX({CR('mes')}),{CF_MONTHS},0),1)"
SUP = [
    ("kgO", "kg Occidente por contenedor", f"=IFERROR(AVERAGE({CR('kgO')}),0)", "kg", "kg"),
    ("kgC", "kg Centro por contenedor", f"=IFERROR(AVERAGE({CR('kgC')}),0)", "kg", "kg"),
    ("kgR", "kg Oriente por contenedor", f"=IFERROR(AVERAGE({CR('kgR')}),0)", "kg", "kg"),
    ("n", "Contenedores por mes", f"=MAX(1,IFERROR(COUNT({CR('n')})/COUNTIF({RS('B')},\">0\"),1))", "num", "cont./mes"),
    ("tO", "Tarifa Occidente", "=" + lastp("B"), "usdkg", "USD/kg"),
    ("tC", "Tarifa Centro", "=" + lastp("C"), "usdkg", "USD/kg"),
    ("tR", "Tarifa Oriente", "=" + lastp("D"), "usdkg", "USD/kg"),
    ("pd", "Precio del diésel", "=" + lastp("E"), "usdkg", "USD/L"),
    ("cita", "Cita del puerto", "=" + lastp("F"), "usd", "USD/viaje"),
    ("dieta", "Dieta del chofer", "=" + lastp("G"), "usd", "USD/viaje"),
    ("jor", "Pago por jornal – transitaria", "=" + lastp("H"), "usd", "USD"),
    ("jorCT", "Pago por jornal – centro transitorio", "=" + lastp("I"), "usd", "USD"),
    ("serv", "Servicio de la transitaria (promedio ENCI / A.V)", f"=({lastp('J')}+{lastp('K')})/2", "usd", "USD/cont."),
    ("cpct", "Comisión variable (% del ingreso)", "=" + lastp("L"), "pct", "%"),
    ("ckg", "Comisión variable (USD/kg)", "=" + lastp("M"), "usdkg", "USD/kg"),
    ("desg", "Desgaste del vehículo propio", "=" + lastp("N"), "usdkg", "USD/km"),
    ("otros", "Estadía y otros directos medios por contenedor",
     f"=IFERROR(AVERAGE({CR('estad')})+AVERAGE({CR('otrG')}),0)", "usd", "USD/cont."),
    ("desv", "Desvío del combustible frente a la norma",
     f"=IFERROR(SUM({CR('litR')})/SUM({CR('litS')})-1,0)", "pct", "%"),
    ("fG", "Fijos generales mensuales", f"=INDEX({cf_row('GENERAL')},{REFMES})", "usd0", "USD/mes"),
    ("fO", "Fijos Occidente mensuales", f"=INDEX({cf_row('OCC')},{REFMES})", "usd0", "USD/mes"),
    ("fC", "Fijos Centro mensuales", f"=INDEX({cf_row('CEN')},{REFMES})", "usd0", "USD/mes"),
    ("fR", "Fijos Oriente mensuales", f"=INDEX({cf_row('ORI')},{REFMES})", "usd0", "USD/mes"),
]
SV = {}
for i, (k, lab, f, fm, unit) in enumerate(SUP):
    rr = 5 + i
    SV[k] = f"$E${rr}"
    put(sm, f"B{rr}", lab, border=True)
    put(sm, f"C{rr}", f, F_LINK, fm, border=True)
    inp(sm, f"D{rr}", None, fm)
    put(sm, f"E{rr}", f'=IF(D{rr}="",C{rr},D{rr})', F_BOLD, fm, border=True)
    put(sm, f"F{rr}", unit, border=True)
dv_num(sm, f"D5:D{4 + len(SUP)}")
put(sm, f"B{5 + len(SUP)}", "Fijos y comisiones: se toman del último mes con contenedores y de la última vigencia de PRECIOS.", F_SUB)

# cálculo del contenedor tipo
B0 = 5 + len(SUP) + 3
put(sm, f"B{B0 - 1}", "Resultado de un contenedor tipo y del mes", F_H2)
for c, t in zip("BCDEF", ["Concepto", "Occidente", "Centro", "Oriente", "TOTAL"]):
    hdr(sm, f"{c}{B0}", t)
v = SV
PCOM = f'IF(Base_Comision="% DEL INGRESO",{v["cpct"]},0)'
KCOM = f'IF(Base_Comision="% DEL INGRESO",0,{v["ckg"]})'
REVP = "(Pct_Reclamaciones+Pct_Banco+Imp_Ventas_Serv+Contrib_Territorial)"
kgT = f'({v["kgO"]}+{v["kgC"]}+{v["kgR"]})'
FUEL = f'(1+{v["desv"]})*{v["pd"]}'
ROWS = {}
SIMR = [
    ("kg", "kg", [v["kgO"], v["kgC"], v["kgR"]], "kg"),
    ("ing", "Ingresos", [f'{v["kgO"]}*{v["tO"]}', f'{v["kgC"]}*{v["tC"]}', f'{v["kgR"]}*{v["tR"]}'], "usd0"),
    ("dist", "Distribución (combustible y desagrupe regional)", [
        f'IF({v["kgO"]}=0,0,ROUNDUP({v["kgO"]}/Cap_Camion_OCC,0))*L_Dist_OCC*{FUEL}',
        f'IF({v["kgC"]}=0,0,ROUNDUP({v["kgC"]}/Cap_Camion_CEN,0))*(L_Tras_CEN+L_Dist_CEN)*{FUEL}'
        f'+IF({v["kgC"]}=0,0,ROUNDUP({v["kgC"]}/Rend_Desagrupe_CT,0))*{v["jorCT"]}',
        f'IF({v["kgR"]}=0,0,ROUNDUP({v["kgR"]}/Cap_Camion_ORI,0))*(L_Tras_ORI+L_Dist_ORI)*{FUEL}'
        f'+IF({v["kgR"]}=0,0,ROUNDUP({v["kgR"]}/Rend_Desagrupe_CT,0))*{v["jorCT"]}'], "usd0"),
    ("com", "Comisión variable", None, "usd0"),
    ("pct", "Reclamaciones, banca y tributos sobre ingresos", None, "usd0"),
    ("fij", "Fijos regionales por contenedor", [f'{v["fO"]}/{v["n"]}', f'{v["fC"]}/{v["n"]}', f'{v["fR"]}/{v["n"]}'], "usd0"),
    ("sha", "Puerto, transitaria, otros y fijos generales (según kg)", None, "usd0"),
    ("uai", "UTILIDAD ANTES DE IMPUESTO POR CONTENEDOR", None, "usd0"),
    ("mg", "Margen %", None, "pct"),
    ("ukg", "Utilidad por kg", None, "usdkg"),
    ("tmin", "TARIFA MÍNIMA DE EQUILIBRIO (USD/kg)", None, "usdkg"),
]
for i, (k, lab, *_x) in enumerate(SIMR):
    ROWS[k] = B0 + 1 + i
R_ = ROWS
SHARED_TOT = (f'(Std_Viajes_Puerto*({v["cita"]}+{v["dieta"]}+Litros_Puerto_Viaje*{FUEL}+Km_Puerto_Viaje*{v["desg"]})'
              f'+IF({kgT}=0,0,ROUNDUP({kgT}/Rend_Desagrupe,0))*{v["jor"]}+{v["serv"]}+{v["otros"]}+{v["fG"]}/{v["n"]})')
for i, (k, lab, fs, fm) in enumerate(SIMR):
    rr = R_[k]
    strong = lab.isupper()
    put(sm, f"B{rr}", lab, F_BOLD if strong else F_BASE, border=True)
    for j, c in enumerate("CDE"):
        if fs:
            f = "=" + fs[j]
        elif k == "com":
            f = f"={c}{R_['ing']}*{PCOM}+{c}{R_['kg']}*{KCOM}"
        elif k == "pct":
            f = f"={c}{R_['ing']}*{REVP}"
        elif k == "sha":
            f = f"=IF({kgT}=0,0,{SHARED_TOT}*{c}{R_['kg']}/{kgT})"
        elif k == "uai":
            f = f"={c}{R_['ing']}-{c}{R_['dist']}-{c}{R_['com']}-{c}{R_['pct']}-{c}{R_['fij']}-{c}{R_['sha']}"
        elif k == "mg":
            f = f"=IF({c}{R_['ing']}=0,0,{c}{R_['uai']}/{c}{R_['ing']})"
        elif k == "ukg":
            f = f"=IF({c}{R_['kg']}=0,0,{c}{R_['uai']}/{c}{R_['kg']})"
        elif k == "tmin":
            f = (f'=IF({c}{R_["kg"]}=0,"s/d",({c}{R_["dist"]}+{c}{R_["kg"]}*{KCOM}+{c}{R_["fij"]}+{c}{R_["sha"]})'
                 f'/({c}{R_["kg"]}*(1-{REVP}-{PCOM})))')
        put(sm, f"{c}{rr}", f, F_BOLD if strong else F_BASE, fm, border=True)
    if k in ("mg",):
        tf = f"=IF(F{R_['ing']}=0,0,F{R_['uai']}/F{R_['ing']})"
    elif k == "ukg":
        tf = f"=IF(F{R_['kg']}=0,0,F{R_['uai']}/F{R_['kg']})"
    elif k == "tmin":
        tf = (f'=IF(F{R_["kg"]}=0,"s/d",(F{R_["dist"]}+F{R_["kg"]}*{KCOM}+F{R_["fij"]}+F{R_["sha"]})'
              f'/(F{R_["kg"]}*(1-{REVP}-{PCOM})))')
    else:
        tf = f"=SUM(C{rr}:E{rr})"
    put(sm, f"F{rr}", tf, F_BOLD, fm, FILL_TOTAL, border=True)
put(sm, f"B{R_['tmin'] + 1}", "Tarifa mínima: precio por kg con el que la región cubre exactamente sus costos y su parte de los compartidos. "
                              "En TOTAL, tarifa única equivalente.", F_SUB)
M0 = R_["tmin"] + 3
put(sm, f"B{M0}", "Utilidad antes de impuesto del MES", F_BOLD, border=True)
put(sm, f"F{M0}", f"=F{R_['uai']}*{v['n']}", F_BOLD, "usd0", FILL_TOTAL, border=True)
put(sm, f"B{M0 + 1}", "Impuesto sobre utilidades estimado del mes", border=True)
put(sm, f"F{M0 + 1}", f"=MAX(0,F{M0})*Imp_Utilidades", fmt="usd0", border=True)
put(sm, f"B{M0 + 2}", "Utilidad neta estimada del mes", F_BOLD, border=True)
put(sm, f"F{M0 + 2}", f"=F{M0}-F{M0 + 1}", F_BOLD, "usd0", FILL_TOTAL, border=True)
# contribución por contenedor (antes de fijos) y punto de equilibrio
put(sm, f"B{M0 + 3}", "Contribución por contenedor (antes de los costos fijos)", border=True)
CONTRIB = (f"($F${R_['ing']}-$F${R_['dist']}-$F${R_['com']}-$F${R_['pct']}"
           f"-({SHARED_TOT}-{v['fG']}/{v['n']}))")
put(sm, f"F{M0 + 3}", f"={CONTRIB}", fmt="usd0", border=True)
FIJT = f"({v['fG']}+{v['fO']}+{v['fC']}+{v['fR']})"
put(sm, f"B{M0 + 4}", "PUNTO DE EQUILIBRIO (contenedores por mes)", F_BOLD, border=True)
put(sm, f"F{M0 + 4}", f'=IF(F{M0 + 3}<=0,"No cubre",{FIJT}/F{M0 + 3})', F_BOLD, "num", FILL_TOTAL, border=True)

# sensibilidad 1: tarifa x diésel
S0 = M0 + 7
put(sm, f"B{S0 - 1}", "Sensibilidad: utilidad mensual antes de impuesto según la tarifa y el precio del diésel", F_H2)
put(sm, f"B{S0}", "Variación de las tarifas ↓ / precio del diésel (USD/L) →", F_BOLD, fill=FILL_GRP, border=True)
DIESEL = [1.50, 1.75, 2.00, 2.25, 2.50, 2.75, 3.00]
FACT = [-0.20, -0.15, -0.10, -0.05, 0, 0.05, 0.10, 0.15, 0.20]
for j, p in enumerate(DIESEL):
    inp(sm, f"{CL(3 + j)}{S0}", p, "usdkg")
# auxiliares (fila oculta de cálculo)
AUX = S0 + len(FACT) + 2
put(sm, f"B{AUX}", "Auxiliares de la tabla (no editar)", F_SUB)
put(sm, f"B{AUX + 1}", "Ingreso base por contenedor", border=True)
put(sm, f"C{AUX + 1}", f"=F{R_['ing']}", fmt="usd0", border=True)
put(sm, f"B{AUX + 2}", "Litros por contenedor (con desvío)", border=True)
LT = (f"(Std_Viajes_Puerto*Litros_Puerto_Viaje"
      f"+IF({v['kgO']}=0,0,ROUNDUP({v['kgO']}/Cap_Camion_OCC,0))*L_Dist_OCC"
      f"+IF({v['kgC']}=0,0,ROUNDUP({v['kgC']}/Cap_Camion_CEN,0))*(L_Tras_CEN+L_Dist_CEN)"
      f"+IF({v['kgR']}=0,0,ROUNDUP({v['kgR']}/Cap_Camion_ORI,0))*(L_Tras_ORI+L_Dist_ORI))*(1+{v['desv']})")
put(sm, f"C{AUX + 2}", "=" + LT, fmt="num", border=True)
put(sm, f"B{AUX + 3}", "Costos variables sin combustible ni % sobre ingreso", border=True)
put(sm, f"C{AUX + 3}",
    f"=(F{R_['dist']}+F{R_['sha']}-{v['fG']}/{v['n']})-(C{AUX + 2}*{v['pd']})+{kgT}*{KCOM}",
    fmt="usd0", border=True)
put(sm, f"B{AUX + 4}", "Costos fijos totales del mes", border=True)
put(sm, f"C{AUX + 4}", f"={FIJT}", fmt="usd0", border=True)
for i, fct in enumerate(FACT):
    rr = S0 + 1 + i
    inp(sm, f"B{rr}", fct, '+0%;-0%;0%')
    for j in range(len(DIESEL)):
        Lc = CL(3 + j)
        put(sm, f"{Lc}{rr}",
            f"={v['n']}*($C${AUX + 1}*(1+$B{rr})*(1-{REVP}-{PCOM})-$C${AUX + 2}*{Lc}${S0}-$C${AUX + 3})-$C${AUX + 4}",
            fmt="usd0", border=True)
sm.conditional_formatting.add(
    f"C{S0 + 1}:{CL(2 + len(DIESEL))}{S0 + len(FACT)}",
    FormulaRule(formula=[f"C{S0 + 1}<0"], fill=PatternFill("solid", fgColor="FCE4D6")))
sm.conditional_formatting.add(
    f"C{S0 + 1}:{CL(2 + len(DIESEL))}{S0 + len(FACT)}",
    FormulaRule(formula=[f"C{S0 + 1}>=0"], fill=PatternFill("solid", fgColor="E2EFDA")))
# sensibilidad 2: contenedores por mes
V0 = AUX + 7
put(sm, f"B{V0 - 1}", "Sensibilidad: utilidad mensual según la cantidad de contenedores", F_H2)
hdr(sm, f"B{V0}", "Contenedores por mes")
hdr(sm, f"C{V0}", "Utilidad del mes")
hdr(sm, f"D{V0}", "Utilidad por contenedor")
for i in range(10):
    rr = V0 + 1 + i
    inp(sm, f"B{rr}", i + 1, "int")
    put(sm, f"C{rr}", f"=B{rr}*({CONTRIB})-{FIJT}", fmt="usd0", border=True)
    put(sm, f"D{rr}", f"=C{rr}/B{rr}", fmt="usd0", border=True)
protect(sm)

# ================================================================ textos: GUIA / ANALISIS / INICIO


def write_text(ws, blocks, start=4):
    ws.column_dimensions["A"].width = 3
    widths(ws, {"B": 34, "C": 48, "D": 40, "E": 30, "F": 22})
    r = start
    num = 0
    for kind, val in blocks:
        if kind == "h1":
            put(ws, f"B{r}", val, F_TITLE)
            r += 2
        elif kind == "h2":
            put(ws, f"B{r}", val, F_H2)
            num = 0
            r += 1
        elif kind in ("p", "b", "n"):
            if kind == "n":
                num += 1
            txt = val if kind == "p" else (f"•  {val}" if kind == "b" else f"{num}.  {val}")
            ws.merge_cells(f"B{r}:F{r}")
            put(ws, f"B{r}", txt, align=WRAP)
            ws.row_dimensions[r].height = max(15, 15 * (len(txt) // 150 + 1))
            r += 1
            if kind == "p":
                r += 1
        elif kind == "t":
            ncols = len(val[0])
            cols = "BCDEF"[:ncols]
            for i, rowv in enumerate(val):
                h = 15
                for c, x in zip(cols, rowv):
                    if i == 0:
                        hdr(ws, f"{c}{r}", x)
                    else:
                        put(ws, f"{c}{r}", x, align=WRAP, border=True)
                    w = ws.column_dimensions[c].width
                    h = max(h, 14 * (len(str(x)) // max(int(w * 1.1), 1) + 1))
                ws.row_dimensions[r].height = h
                r += 1
            r += 1
    return r


gu = sheet("GUIA_USO", "FFC000")
write_text(gu, GUIA, 2)
protect(gu)
an = sheet("ANALISIS", "FFC000")
write_text(an, ANALISIS, 2)
put(an, "B1", "Fuentes de tributos y combustible: ver README (enlaces). Valores ILUSTRATIVOS marcados en PARAMETROS y PRECIOS.", F_SUB)
protect(an)

ini = sheet("INICIO", "1F3864")
wb.move_sheet("INICIO", offset=-(len(wb.sheetnames) - 1))
ini.column_dimensions["A"].width = 3
widths(ini, {"B": 38, "C": 70, "D": 28})
put(ini, "B2", "ANÁLISIS DE RENTABILIDAD – NEGOCIO DE PAQUETERÍA", F_TITLE)
put(ini, "B3", '=Empresa&"  ·  Modelo por contenedor, por región y por mes  ·  Importes en USD"', F_SUB)
put(ini, "B5", "Flujo modelado", F_H2)
ini.merge_cells("B6:D6")
put(ini, "B6", "1. PUERTO (vehículo propio: cita, combustible, dieta, desgaste)  →  2. TRANSITARIA ENCI / A.V (desagrupe por jornal)  →  "
               "3. DISTRIBUCIÓN: Occidente directa · Centro y Oriente con traslado, centro transitorio y distribución  →  "
               "4. COMERCIALIZACIÓN (21 gestores: fijo + variable)", align=WRAP)
ini.row_dimensions[6].height = 45
put(ini, "B8", "Hojas del libro", F_H2)
hdr(ini, "B9", "Hoja")
hdr(ini, "C9", "Contenido")
hdr(ini, "D9", "Cuándo se usa")
HOJAS = [
    ("GUIA_USO", "Resumen de la guía (la guía completa con imágenes es el PDF que acompaña al libro).", "Antes de empezar"),
    ("ANALISIS", "Diagnóstico, metodología, clasificación de costos y propuestas de mejora.", "Para decidir"),
    ("PARAMETROS", "Normas técnicas, porcentajes, tributos y métodos de cálculo.", "Una vez / si cambia una norma"),
    ("PRECIOS", "Tarifas y precios con fecha de vigencia (historial auditable).", "Una vez / si cambia un precio"),
    ("ACTIVOS", "Vehículos y equipos: depreciación, mantenimiento y seguros.", "Una vez / al comprar o vender"),
    ("COSTOS_FIJOS", "Costos fijos mensuales por centro de costo y tributos sobre salarios.", "Una vez / cierre de mes"),
    ("COMERCIAL", "21 gestores: pago fijo, liquidación variable y conciliación mensual.", "Cierre de mes"),
    ("ENTRADA", "Registro de contenedores: la hoja donde se trabaja a diario.", "Con cada contenedor"),
    ("CALCULO", "Cálculo automático completo por contenedor y por región.", "Solo consulta"),
    ("RESUMEN_MES", "Rentabilidad mensual, punto de equilibrio e impuesto anual estimado.", "Cierre de mes"),
    ("DASHBOARD", "Panel de indicadores y gráficos del período elegido.", "Para informar"),
    ("SIMULADOR", "Escenarios: tarifa mínima por región, sensibilidad a la tarifa, al diésel y al volumen.", "Para decidir"),
    ("LISTAS", "Listas de los desplegables (técnica).", "No se toca"),
]
for i, (h, d, q) in enumerate(HOJAS):
    rr = 10 + i
    c = put(ini, f"B{rr}", h, Font(name=FONT, size=10, color="0563C1", underline="single"), border=True)
    c.hyperlink = f"#'{h}'!A1"
    put(ini, f"C{rr}", d, border=True, align=WRAP)
    put(ini, f"D{rr}", q, border=True)
RU = 10 + len(HOJAS) + 1
put(ini, f"B{RU}", "Su rutina de trabajo (una sola persona lleva el libro)", F_H2)
RUTINA = [
    ("Una sola vez, al empezar", "PARAMETROS → PRECIOS → ACTIVOS → COSTOS_FIJOS → COMERCIAL. Después, borrar los datos de ejemplo."),
    ("Con cada contenedor", "ENTRADA: 1) al llegar, ID, fecha, transitaria y kg por región; 2) al terminar la distribución, datos reales y Estado «Cerrado»."),
    ("Cuando cambia un precio", "PRECIOS: agregar una fila nueva con la fecha desde la que rige. Nunca sobrescribir la anterior."),
    ("Cada fin de mes", "Revisar alertas en INICIO → gastos reales en COSTOS_FIJOS → liquidación en COMERCIAL → RESUMEN_MES y DASHBOARD → guardar una copia con fecha."),
]
for i, (cu, qu) in enumerate(RUTINA):
    rr = RU + 1 + i
    put(ini, f"B{rr}", cu, F_BOLD, border=True)
    ini.merge_cells(f"C{rr}:D{rr}")
    put(ini, f"C{rr}", qu, border=True, align=WRAP)
    ini.row_dimensions[rr].height = 28
LG = RU + len(RUTINA) + 2
put(ini, f"B{LG}", "Leyenda de colores", F_H2)
put(ini, f"B{LG + 1}", "1.234", F_INPUT, fmt="#,##0", fill=FILL_INPUT, border=True)
put(ini, f"C{LG + 1}", "Dato de entrada: se puede escribir (letra azul, fondo amarillo).")
put(ini, f"B{LG + 2}", 1234, F_BASE, "#,##0", border=True)
put(ini, f"C{LG + 2}", "Fórmula: no escribir (letra negra).")
put(ini, f"B{LG + 3}", 1234, F_LINK, "#,##0", border=True)
put(ini, f"C{LG + 3}", "Enlace a otra hoja (letra verde).")
ST = LG + 5
put(ini, f"B{ST}", "Estado del libro (controles automáticos)", F_H2)
CHECKS = [
    ("Contenedores registrados", f'=COUNTIF(ENTRADA!$B${DATA0}:$B${DATAN},"?*")', "int"),
    ("Contenedores con alertas", f'=COUNTIF({CR("alert")},"?*")', "int"),
    ("Contenedores abiertos (sin cerrar)", f'=COUNTIF({CR("estado")},"Abierto")', "int"),
    ("Cuadre de regiones (CALCULO)", f'=IF(AND(MAX({CR("chk")})<0.01,MIN({CR("chk")})>-0.01),"OK","REVISAR")', None),
    ("Cuadre mensual (RESUMEN_MES)", f'=IF(AND(MAX({RS("AA")})<0.01,MIN({RS("AA")})>-0.01),"OK","REVISAR")', None),
    ("Conciliación de comercialización", f"=COMERCIAL!G{CO_CONC}", None),
    ("Utilidad antes de impuesto acumulada", f"=RESUMEN_MES!O{RT}", "usd0"),
]
for i, (lab, f, fm) in enumerate(CHECKS):
    rr = ST + 1 + i
    put(ini, f"B{rr}", lab, border=True)
    put(ini, f"C{rr}", f, F_BOLD, fm, border=True, align=Alignment(horizontal="left"))
ini.conditional_formatting.add(f"C{ST + 4}:C{ST + 6}", FormulaRule(formula=[f'LEFT(C{ST + 4},2)="OK"'],
                                                                  font=Font(name=FONT, bold=True, color="008000")))
ini.conditional_formatting.add(f"C{ST + 4}:C{ST + 6}", FormulaRule(formula=[f'LEFT(C{ST + 4},3)="REV"'],
                                                                  font=Font(name=FONT, bold=True, color="C00000")))
ini.conditional_formatting.add(f"C{ST + 6}", FormulaRule(formula=[f'C{ST + 6}="CUADRA"'],
                                                         font=Font(name=FONT, bold=True, color="008000")))
put(ini, f"B{ST + 10}", "Los datos que contiene el libro son de EJEMPLO para mostrar su funcionamiento. Sustitúyalos siguiendo la guía en PDF (capítulo «Preparar el libro con sus datos»).",
    Font(name=FONT, size=10, bold=True, color="C00000"))
protect(ini)

# orden de hojas
ORDER = ["INICIO", "GUIA_USO", "ANALISIS", "DASHBOARD", "ENTRADA", "PARAMETROS", "PRECIOS", "ACTIVOS",
         "COSTOS_FIJOS", "COMERCIAL", "CALCULO", "RESUMEN_MES", "SIMULADOR", "LISTAS"]
wb._sheets = [wb[n] for n in ORDER]
wb.active = 0
wb.calculation.fullCalcOnLoad = True
for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print("OK", OUT)
