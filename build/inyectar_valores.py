# -*- coding: utf-8 -*-
"""Copia los valores calculados por LibreOffice (archivo recalculado) dentro del
libro original de openpyxl, sin tocar fórmulas, validaciones ni protecciones.

Así el archivo entregado muestra resultados en cualquier visor y Excel los
vuelve a calcular al abrir (fullCalcOnLoad).

Uso: python inyectar_valores.py original.xlsx recalculado.xlsx salida.xlsx
"""
import datetime as dt
import sys
import zipfile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
orig, recalc, out = sys.argv[1:4]

vals = load_workbook(recalc, data_only=True)

with zipfile.ZipFile(orig) as z:
    wbx = etree.fromstring(z.read("xl/workbook.xml"))
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels}
    sheet_xml = {}
    for s in wbx.iter(f"{{{NS}}}sheet"):
        t = target[s.get(f"{{{RNS}}}id")].lstrip("/")
        sheet_xml["xl/" + t if not t.startswith("xl/") else t] = s.get("name")

    filled = 0
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in z.infolist():
            data = z.read(item.filename)
            if item.filename in sheet_xml:
                ws = vals[sheet_xml[item.filename]]
                root = etree.fromstring(data)
                for c in root.iter(f"{{{NS}}}c"):
                    f = c.find(f"{{{NS}}}f")
                    if f is None:
                        continue
                    v = ws[c.get("r")].value
                    old = c.find(f"{{{NS}}}v")
                    if old is not None:
                        c.remove(old)
                    if v is None:
                        continue
                    ve = etree.SubElement(c, f"{{{NS}}}v")
                    if isinstance(v, bool):
                        c.set("t", "b")
                        ve.text = "1" if v else "0"
                    elif isinstance(v, (int, float)):
                        c.attrib.pop("t", None)
                        ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                    elif isinstance(v, (dt.datetime, dt.date)):
                        c.attrib.pop("t", None)
                        ve.text = repr(float(to_excel(v)))
                    else:
                        c.set("t", "str")
                        ve.text = str(v)
                    filled += 1
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zo.writestr(item, data)
print(f"valores inyectados: {filled}")
